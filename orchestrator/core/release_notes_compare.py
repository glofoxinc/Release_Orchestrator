"""Compare generated Release Notes DOCX against a reference file."""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any

from docx import Document

from generators.release_notes_parity import document_fingerprint, parity_gaps


V4_SECTION_MARKERS = (
    "Release Notes",
    "Release Summary",
    "Key Highlights",
    "Reports / Dashboards Delivered",
    "Ticket Details:",
    "Testing Evidence",
    "Tickets Excluded from Deployment",
    "Notes:",
)


def extract_section_headings(doc: Document) -> list[str]:
    """Return non-empty paragraph lines that look like section headings."""
    headings: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        first_line = text.split("\n")[0].strip()
        if first_line in V4_SECTION_MARKERS or first_line.endswith(":"):
            headings.append(first_line)
        elif first_line.startswith("Version:"):
            headings.append(first_line[:80])
    return headings


def _doc_contains_marker(doc: Document, marker: str) -> bool:
    blob = "\n".join(p.text for p in doc.paragraphs)
    return marker in blob


def compare_docx_paths(reference_path: str | Path, generated_path: str | Path) -> dict[str, Any]:
    """Structural compare: section markers + document fingerprint gaps."""
    ref_path = Path(reference_path)
    gen_path = Path(generated_path)
    if not ref_path.is_file():
        return {"status": "error", "error": f"Reference not found: {ref_path}"}
    if not gen_path.is_file():
        return {"status": "error", "error": f"Generated not found: {gen_path}"}

    ref_doc = Document(str(ref_path))
    gen_doc = Document(str(gen_path))
    ref_headings = extract_section_headings(ref_doc)
    gen_headings = extract_section_headings(gen_doc)
    missing_sections = [m for m in V4_SECTION_MARKERS if not _doc_contains_marker(gen_doc, m)]
    ref_fp = document_fingerprint(ref_doc)
    gen_fp = document_fingerprint(gen_doc)
    gaps = parity_gaps(ref_fp, gen_fp)
    ok = not missing_sections and len(gaps) <= 14
    return {
        "status": "ok" if ok else "fail",
        "reference": str(ref_path),
        "generated": str(gen_path),
        "reference_table_count": len(ref_doc.tables),
        "generated_table_count": len(gen_doc.tables),
        "missing_v4_sections": missing_sections,
        "fingerprint_gaps": gaps,
        "fingerprint_gap_count": len(gaps),
        "reference_headings": ref_headings[:20],
        "generated_headings": gen_headings[:20],
    }


def compare_generated_bytes(reference_path: str | Path, docx_bytes: bytes) -> dict[str, Any]:
    """Compare reference file to in-memory DOCX bytes."""
    gen_doc = Document(io.BytesIO(docx_bytes))
    ref_path = Path(reference_path)
    if not ref_path.is_file():
        return {"status": "error", "error": f"Reference not found: {ref_path}"}
    ref_doc = Document(str(ref_path))
    missing_sections = [
        m for m in V4_SECTION_MARKERS
        if not _doc_contains_marker(gen_doc, m)
    ]
    gaps = parity_gaps(document_fingerprint(ref_doc), document_fingerprint(gen_doc))
    ok = not missing_sections and len(gaps) <= 14
    return {
        "status": "ok" if ok else "fail",
        "reference": str(ref_path),
        "missing_v4_sections": missing_sections,
        "fingerprint_gaps": gaps,
        "fingerprint_gap_count": len(gaps),
    }
