"""Structural fingerprint helpers for reference DOCX parity tests."""

from __future__ import annotations

from docx import Document
from docx.oxml.ns import qn

from generators.release_notes_template_defaults import CHANGE_COLUMN_COUNT

TABLE_STYLE_ID_TO_NAME = {
    "GridTable4-Accent1": "Grid Table 4 Accent 1",
    "TableGrid": "Table Grid",
}


def table_style_name(table) -> str | None:
    """Resolve table style from w:tblStyle (reference parity) or python-docx style."""
    tbl_pr = table._tbl.tblPr
    if tbl_pr is not None:
        style_el = tbl_pr.find(qn("w:tblStyle"))
        if style_el is not None:
            style_id = style_el.get(qn("w:val"))
            return TABLE_STYLE_ID_TO_NAME.get(style_id, style_id)
    return table.style.name if table.style else None


def table_grid_widths(table) -> tuple[int, ...]:
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is None:
        return ()
    return tuple(int(col.get(qn("w:w"))) for col in grid.findall(qn("w:gridCol")))


def table_total_width(table) -> int | None:
    tbl_pr = table._tbl.tblPr
    if tbl_pr is None:
        return None
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        return None
    value = tbl_w.get(qn("w:w"))
    if value in (None, "0"):
        return 0
    return int(value)


def table_layout(table) -> str | None:
    tbl_pr = table._tbl.tblPr
    if tbl_pr is None:
        return None
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        return None
    return layout.get(qn("w:type"))


def table_has_borders(table) -> bool:
    tbl_pr = table._tbl.tblPr
    if tbl_pr is None:
        return False
    return tbl_pr.find(qn("w:tblBorders")) is not None


def cell_has_borders(cell) -> bool:
    tc_pr = cell._tc.tcPr
    if tc_pr is None:
        return False
    return tc_pr.find(qn("w:tcBorders")) is not None


def paragraph_spacing(paragraph) -> dict[str, str] | None:
    p_pr = paragraph._p.pPr
    if p_pr is None:
        return None
    sp = p_pr.find(qn("w:spacing"))
    if sp is None:
        return None
    out: dict[str, str] = {}
    for key in ("before", "after", "line", "lineRule"):
        value = sp.get(qn(f"w:{key}"))
        if value is not None:
            out[key] = value
    return out or None


def style_color_hex(style) -> str | None:
    if style.font.color and style.font.color.rgb:
        return str(style.font.color.rgb)
    return None


def run_color(run) -> str | None:
    if run.font.color and run.font.color.rgb:
        return str(run.font.color.rgb)
    return None


def paragraph_shading(paragraph) -> str | None:
    p_pr = paragraph._p.pPr
    if p_pr is None:
        return None
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        return None
    return shd.get(qn("w:fill"))


def header_cell_fingerprint(cell) -> dict:
    tc_pr = cell._tc.tcPr
    shd = tc_pr.find(qn("w:shd")) if tc_pr is not None else None
    tc_w = tc_pr.find(qn("w:tcW")) if tc_pr is not None else None
    run = cell.paragraphs[0].runs[0] if cell.paragraphs and cell.paragraphs[0].runs else None
    font = None
    if run is not None:
        font = (
            run.font.name,
            run.font.size.pt if run.font.size else None,
            run_color(run),
            run.bold,
        )
    return {
        "text": cell.text.strip(),
        "fill": shd.get(qn("w:fill")) if shd is not None else None,
        "border": cell_has_borders(cell),
        "tcW": tc_w.get(qn("w:w")) if tc_w is not None else None,
        "nowrap": tc_pr.find(qn("w:noWrap")) is not None if tc_pr is not None else False,
        "align": cell.paragraphs[0].alignment if cell.paragraphs else None,
        "font": font,
        "runs_empty": not bool(cell.paragraphs and cell.paragraphs[0].runs),
    }


def _deploy_table(doc: Document):
    return next(
        t for t in doc.tables if any("Deployment Owner" in c.text for c in t.rows[0].cells)
    )


def _change_table(doc: Document):
    """Ticket Details table — 5-column layout, with legacy 10-column fallback."""
    candidates = []
    for table in doc.tables:
        if not table.rows:
            continue
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        if headers and headers[0] in {"Ticket", "S.No"} and any(
            token in headers for token in ("Ticket Link", "Ticket#", "Link")
        ):
            candidates.append((len(table.columns), table, headers))

    for col_count, table, headers in candidates:
        if col_count == CHANGE_COLUMN_COUNT:
            return table

    for col_count, table, headers in candidates:
        if col_count == 4 and "Ticket Link" in headers and "Assignee" not in headers:
            return table

    for col_count, table, headers in sorted(candidates, key=lambda item: item[0], reverse=True):
        if col_count >= 10 and "Description" in headers:
            return table

    raise ValueError("No Ticket Details table found in document")


def _excluded_table(doc: Document):
    for table in doc.tables:
        if len(table.columns) != 8:
            continue
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        if headers and headers[0] == "S.No" and "Ticket#" in headers:
            return table
    if len(doc.tables) >= 3:
        return doc.tables[2]
    raise ValueError("No excluded tickets table found in document")


def document_fingerprint(doc: Document) -> dict:
    shaded = [
        paragraph.text.strip().split("\n")[0][:50]
        for paragraph in doc.paragraphs
        if paragraph.text.strip() and paragraph_shading(paragraph) == "F7F7F7"
    ]
    fp: dict = {
        "title_runs": [
            (run.text, run.font.name, run.font.size.pt if run.font.size else None, run_color(run), run.bold)
            for run in doc.paragraphs[0].runs
        ],
        "version_runs": [
            (run.text, run.bold, run_color(run))
            for run in doc.paragraphs[1].runs
        ],
        "heading1_style_color": style_color_hex(doc.styles["Heading 1"]),
        "heading2_style_color": style_color_hex(doc.styles["Heading 2"]),
        "heading3_style_color": style_color_hex(doc.styles["Heading 3"]),
        "shaded_count": len(shaded),
        "shaded_headings": shaded,
    }
    kh = next((p for p in doc.paragraphs if p.text.strip() == "Key Highlights"), None)
    if kh is not None:
        fp["key_highlights_spacing"] = paragraph_spacing(kh)
    named_tables = {
        "deploy": _deploy_table(doc),
        "change": _change_table(doc),
        "excluded": _excluded_table(doc),
    }
    for name, table in named_tables.items():
        fp[f"{name}_tblW"] = table_total_width(table)
        fp[f"{name}_layout"] = table_layout(table)
        fp[f"{name}_tblBorders"] = table_has_borders(table)
        fp[f"{name}_widths"] = table_grid_widths(table)
        fp[f"{name}_style"] = table_style_name(table)
        fp[f"{name}_hdr0"] = header_cell_fingerprint(table.rows[0].cells[0])
        if name == "deploy" and len(table.rows) > 1:
            fp["deploy_owner1"] = header_cell_fingerprint(table.rows[1].cells[1])
            fp["deploy_status1"] = header_cell_fingerprint(table.rows[1].cells[2])
            fp["deploy_idtime1"] = header_cell_fingerprint(table.rows[1].cells[3])
    return fp


def parity_gaps(reference: dict, generated: dict) -> list[str]:
    """Return human-readable gaps between reference and generated fingerprints."""
    gaps: list[str] = []
    keys = sorted(set(reference) | set(generated))
    for key in keys:
        if reference.get(key) != generated.get(key):
            gaps.append(key)
    return gaps
