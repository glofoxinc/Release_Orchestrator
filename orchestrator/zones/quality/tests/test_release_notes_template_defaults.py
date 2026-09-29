"""Ensure workspace release-notes metadata matches org reference template."""

from __future__ import annotations

import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
os.environ.setdefault("JIRA_EMAIL", "test@example.com")
os.environ.setdefault("JIRA_API_TOKEN", "fake")


def test_workspace_release_meta_matches_reference():
    import config

    config.switch_workspace("")
    from generators.release_notes_template_defaults import REFERENCE_META

    assert config.RELEASE_META["director"] == REFERENCE_META["director"]
    assert config.RELEASE_META["svp"] == REFERENCE_META["svp"]
    assert config.RELEASE_META["senior_mgr"] == REFERENCE_META["senior_mgr"]


def test_minimal_docx_includes_all_reference_sections():
    import io

    from docx import Document

    from generators.release_notes_generator import generate_release_notes_docx

    result = generate_release_notes_docx({"version": "v2026.3.07.15"})
    doc = Document(io.BytesIO(result))
    text = "\n".join(p.text for p in doc.paragraphs)

    for phrase in (
        "Testing Evidence:",
        "Tickets Excluded from Deployment",
        "Notes:",
        "Ticket Details:",
    ):
        assert phrase in text

    assert "Sanity Checklist After Production Deployment:" not in text

    assert len(doc.tables) >= 2  # deployment + excluded placeholder (minimal has no change rows)
