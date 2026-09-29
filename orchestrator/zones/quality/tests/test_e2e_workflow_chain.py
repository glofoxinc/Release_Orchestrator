"""End-to-end workflow chain: snapshot → release notes → SharePoint → service ticket."""

from __future__ import annotations

import io
import os
import sys
from unittest.mock import patch

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
os.environ.setdefault("JIRA_EMAIL", "test@example.com")
os.environ.setdefault("JIRA_API_TOKEN", "fake")


@pytest.fixture
def workflow_state():
    return {
        "shared_release_PIC": {"id": "1", "name": "2026.3.07.29"},
        "shared_release_AN": {"id": "2", "name": "2026.3.07.29"},
    }


class TestWorkflowChain:
    def test_full_prerequisite_chain(self, workflow_state):
        import streamlit as st
        from docx import Document

        from generators.release_notes_generator import generate_release_notes_docx
        from release_platform.ui.helpers import create_snapshot, get_snapshot_ticket_keys
        from release_platform.ui.workflow import (
            check_app_prerequisites,
            save_sharepoint_url,
            validate_sharepoint_url,
        )

        issues = [
            {
                "key": "PIC-1",
                "fields": {"summary": "Custom fix", "status": {"name": "Ready"}},
            },
            {
                "key": "AN-1",
                "fields": {"summary": "Analyze fix", "status": {"name": "Ready"}},
            },
        ]

        with patch.object(st, "session_state", workflow_state, create=False):
            assert check_app_prerequisites("release_notes").ok is False

            create_snapshot(issues)
            assert get_snapshot_ticket_keys() == {"PIC-1", "AN-1"}
            assert check_app_prerequisites("release_notes").ok is True

            doc_bytes = generate_release_notes_docx(
                {
                    "version": "v2026.3.07.29",
                    "change_rows": (
                        "1|PIC-1|PIC-1||https://example/browse/PIC-1|"
                        "Custom|Client|Report|Desc|Tester||"
                    ),
                }
            )
            doc = Document(io.BytesIO(doc_bytes))
            text = "\n".join(p.text for p in doc.paragraphs)
            assert "Ticket Details:" in text
            assert len(doc.tables) >= 2

            workflow_state["shared_docx_generated"] = True
            stk = check_app_prerequisites("service_ticket")
            assert stk.ok is False
            assert any("SharePoint" in m for m in stk.missing)

            save_sharepoint_url("https://sharepoint.example/sites/doc.docx")
            assert validate_sharepoint_url("https://sharepoint.example/sites/doc.docx")
            assert check_app_prerequisites("service_ticket").ok is True

            assert check_app_prerequisites("release_deployer").ok is False
            workflow_state["stk_created_key"] = "TCKT-100"
            assert check_app_prerequisites("release_deployer").ok is True

    def test_sharepoint_link_in_tckt_description(self, workflow_state):
        import streamlit as st

        from adf_builder import adf_links
        from release_platform.ui.helpers import release_notes_doc_filename, sync_sharepoint_from_shared
        from release_platform.ui.service_ticket.step3_description import _build_description_adf

        sp_url = "https://sharepoint.example/sites/Release%20Notes%202026.3.07.29.docx"
        workflow_state.update(
            {
                "shared_docx_filename": "Release Notes 2026.3.07.29.docx",
                "shared_docx_generated": True,
                "shared_sharepoint_url": sp_url,
                "stk_release_PIC": {"id": "1", "name": "2026.3.07.29"},
                "stk_release_AN": {"id": "2", "name": "2026.3.07.29"},
                "stk_extracted": [
                    {
                        "type": "Analyze",
                        "client": "Standard",
                        "report": "Visit Analysis",
                        "description": "Fix drill-through",
                    }
                ],
            }
        )

        with patch.object(st, "session_state", workflow_state, create=False):
            sync_sharepoint_from_shared()
            assert workflow_state["stk_sharepoint"] == sp_url
            assert release_notes_doc_filename("stk_") == "Release Notes 2026.3.07.29.docx"

            adf_doc = _build_description_adf(workflow_state["stk_extracted"], sp_url)
            links = adf_links(adf_doc)
            assert len(links) == 1
            label, href = links[0]
            assert label == "Release Notes 2026.3.07.29.docx"
            assert href == sp_url
            assert "2026.3.07.29 + 2026.3.07.29" not in label
            assert "2026.3.07.29.docx + " not in label
