"""Line-by-line DOCX parity vs Release Notes 2026.3.07.15.docx reference."""

from __future__ import annotations

import io
import os
import sys

import pytest
from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from generators.release_notes_generator import generate_release_notes_docx
from generators.release_notes_template_defaults import (
    CHANGE_COLUMN_WIDTHS,
    CHANGE_TABLE_WIDTH,
    DEPLOYMENT_COLUMN_WIDTHS,
    DEPLOYMENT_TABLE_WIDTH,
    EXCLUDED_COLUMN_WIDTHS,
    EXCLUDED_TABLE_WIDTH,
    PAGE_MARGIN_INCHES,
)
from tests.reference_fixture import REFERENCE_AVAILABLE, REFERENCE_PATH
from tests.test_reference_template import _reference_like_data


def _grid_widths(table) -> tuple[int, ...]:
    grid = table._tbl.find(qn("w:tblGrid"))
    assert grid is not None
    return tuple(int(gc.get(qn("w:w"))) for gc in grid.findall(qn("w:gridCol")))


def _tbl_width(table) -> int | None:
    tbl_pr = table._tbl.tblPr
    if tbl_pr is None:
        return None
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        return None
    value = tbl_w.get(qn("w:w"))
    if value in (None, "0"):
        return 0
    return int(value)


def _tbl_layout(table) -> str | None:
    tbl_pr = table._tbl.tblPr
    if tbl_pr is None:
        return None
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        return None
    return layout.get(qn("w:type"))


def _deploy_table(doc):
    return next(
        t for t in doc.tables if any("Deployment Owner" in c.text for c in t.rows[0].cells)
    )


def _change_table_by_header(doc):
    from generators.release_notes_parity import _change_table

    return _change_table(doc)


def _excluded_table(doc):
    return next(
        t for t in doc.tables if t.rows[0].cells[0].text.strip() == "S.No" and len(t.columns) == 8
    )


@pytest.mark.skipif(not REFERENCE_AVAILABLE, reason="Reference DOCX not available")
class TestReferenceDocxLineParity:
    def test_margins_match_reference(self):
        ref = Document(REFERENCE_PATH)
        gen = Document(io.BytesIO(generate_release_notes_docx(_reference_like_data())))
        for doc in (ref, gen):
            section = doc.sections[0]
            assert round(section.left_margin.inches, 2) == PAGE_MARGIN_INCHES
            assert round(section.top_margin.inches, 2) == PAGE_MARGIN_INCHES

    def test_deployment_table_widths_match_template(self):
        """Generated deployment table uses v6 column widths (wider Owner column)."""
        gen = Document(io.BytesIO(generate_release_notes_docx(_reference_like_data())))
        gen_deploy = _deploy_table(gen)
        assert _grid_widths(gen_deploy) == DEPLOYMENT_COLUMN_WIDTHS
        assert _tbl_width(gen_deploy) == DEPLOYMENT_TABLE_WIDTH

    def test_change_table_widths_and_layout_match_reference(self):
        gen = Document(io.BytesIO(generate_release_notes_docx(_reference_like_data())))
        gen_change = _change_table_by_header(gen)
        assert _grid_widths(gen_change) == CHANGE_COLUMN_WIDTHS
        assert _tbl_width(gen_change) == CHANGE_TABLE_WIDTH
        assert _tbl_layout(gen_change) == "fixed"

    def test_change_table_indent_and_look_match_reference(self):
        from generators.release_notes_template_defaults import (
            CHANGE_TBL_IND_TWIPS,
            CHANGE_TBL_LOOK_VAL,
        )

        gen = Document(io.BytesIO(generate_release_notes_docx(_reference_like_data())))
        table = _change_table_by_header(gen)
        tbl_pr = table._tbl.tblPr
        ind = tbl_pr.find(qn("w:tblInd"))
        look = tbl_pr.find(qn("w:tblLook"))
        assert ind is not None
        assert int(ind.get(qn("w:w"))) == CHANGE_TBL_IND_TWIPS
        assert look.get(qn("w:val")) == CHANGE_TBL_LOOK_VAL

    def test_change_table_row_cnf_and_height_match_reference(self):
        gen = Document(io.BytesIO(generate_release_notes_docx(_reference_like_data())))
        from generators.release_notes_template_defaults import CHANGE_DATA_ROW_HEIGHT_TWIPS

        gen_change = _change_table_by_header(gen)
        gen_hdr_tr = gen_change.rows[0]._tr.trPr
        assert gen_hdr_tr is not None
        assert gen_hdr_tr.find(qn("w:cnfStyle")) is not None

        gen_data_tr = gen_change.rows[1]._tr.trPr
        assert gen_data_tr is not None
        gen_h = gen_data_tr.find(qn("w:trHeight"))
        assert gen_h is not None
        assert gen_h.get(qn("w:val")) == str(CHANGE_DATA_ROW_HEIGHT_TWIPS)

        gen_cell_ppr = gen_change.rows[1].cells[0].paragraphs[0]._p.pPr
        assert gen_cell_ppr is not None
        assert gen_cell_ppr.find(qn("w:cnfStyle")) is not None

    def test_excluded_table_widths_match_reference(self):
        ref = Document(REFERENCE_PATH)
        gen = Document(io.BytesIO(generate_release_notes_docx(_reference_like_data())))
        ref_ex = _excluded_table(ref)
        gen_ex = _excluded_table(gen)
        assert _grid_widths(ref_ex) == EXCLUDED_COLUMN_WIDTHS
        assert _grid_widths(gen_ex) == EXCLUDED_COLUMN_WIDTHS
        assert _tbl_width(gen_ex) == EXCLUDED_TABLE_WIDTH
        assert _tbl_layout(gen_ex) == "fixed"

    def test_excluded_table_has_column_headers(self):
        from generators.release_notes_template_defaults import EXCLUDED_HEADER_TEXT, EXCLUDED_HEADERS

        gen = Document(io.BytesIO(generate_release_notes_docx(_reference_like_data())))
        excluded = _excluded_table(gen)
        headers = [c.text.strip() for c in excluded.rows[0].cells]
        assert headers == EXCLUDED_HEADERS
        hdr_cell = excluded.rows[0].cells[0]
        shd = hdr_cell._tc.tcPr.find(qn("w:shd"))
        assert shd is not None
        assert shd.get(qn("w:fill")) == "000000"
        run = hdr_cell.paragraphs[0].runs[0]
        assert str(run.font.color.rgb) == EXCLUDED_HEADER_TEXT
        assert hdr_cell._tc.tcPr.find(qn("w:noWrap")) is not None
        body_fill = excluded.rows[1].cells[0]._tc.tcPr.find(qn("w:shd"))
        assert body_fill is not None
        assert body_fill.get(qn("w:fill")) == "C0C0C0"

    def test_change_table_all_cells_have_borders(self):
        gen = Document(io.BytesIO(generate_release_notes_docx(_reference_like_data())))
        table = _change_table_by_header(gen)
        total = sum(len(row.cells) for row in table.rows)
        bordered = sum(
            1
            for row in table.rows
            for cell in row.cells
            if cell._tc.tcPr is not None and cell._tc.tcPr.find(qn("w:tcBorders")) is not None
        )
        assert bordered == total
