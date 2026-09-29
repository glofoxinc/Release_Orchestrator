"""Tests for tabular Key Highlights release summary."""

from __future__ import annotations

import io
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
os.environ.setdefault("JIRA_EMAIL", "test@example.com")
os.environ.setdefault("JIRA_API_TOKEN", "fake")


def _ticket(
    ticket_type: str,
    client: str,
    report: str,
    description: str,
    ticket: str,
    assignee: str = "",
) -> dict:
    return {
        "type": ticket_type,
        "client": client,
        "report": report,
        "description": description,
        "ticket": ticket,
        "assignee": assignee,
    }


class TestHighlightSummary:
    def test_builds_category_count_table(self):
        from generators.release_notes_summary_tables import (
            build_highlight_summary,
            highlight_summary_tables,
        )

        extracted = [
            _ticket("Analyze", "Standard", "New Dashboard", "New dashboard", "AN-1", "Alice"),
            _ticket(
                "Analyze",
                "Standard",
                "Promo look up",
                "New paginated report",
                "AN-2",
                "Alice",
            ),
            _ticket("Custom", "Jazzercise", "Bookings by Location", "update", "PIC-1", "Bob"),
            _ticket(
                "Custom",
                "Lift",
                "ABC Customize",
                "New paginated report migrated from Looker",
                "PIC-2",
                "Carol",
            ),
        ]
        summary = build_highlight_summary(extracted)
        assert summary.total_tickets == 4
        assert summary.category_counts["New Reports"] >= 2
        assert summary.category_counts["Looker Exits"] >= 1

        tables = highlight_summary_tables(summary)
        assert tables[0].headers == ("Category", "Count")
        rows = {row[0]: row[1] for row in tables[0].rows}
        assert "Total" in rows
        assert rows["Total"] == "4"

    def test_docx_renders_summary_tables_under_release_summary(self):
        from docx import Document

        from generators.release_notes_generator import generate_release_notes_docx
        from generators.release_notes_template_defaults import CATEGORY_SUMMARY_HEADERS

        extracted = [
            _ticket("Analyze", "Standard", "Payouts", "fix", "AN-3", "Alice"),
            _ticket("Custom", "Lift", "ABC Customize", "New paginated report", "PIC-9", "Bob"),
        ]
        doc = Document(
            io.BytesIO(
                generate_release_notes_docx(
                    {
                        "version": "v2026.3.08.12",
                        "extracted": extracted,
                        "include_analyze_section": True,
                        "custom_report_tables": "[Lift]\nLift|ABC Customize|PIC-9|Bob|Paginated report|New",
                        "analyze_report_rows": "Analyze|Payouts|AN-3|Alice|Dashboard|Existing",
                        "custom_clients": "Lift",
                    }
                )
            )
        )
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Key Highlights" in text
        assert "Release Summary" in text
        summary_tables = [
            table
            for table in doc.tables
            if [cell.text.strip() for cell in table.rows[0].cells] == list(CATEGORY_SUMMARY_HEADERS)
        ]
        assert summary_tables
        release_summary_idx = next(
            i for i, p in enumerate(doc.paragraphs) if "Release Summary" in p.text
        )
        key_highlights_idx = next(
            i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "Key Highlights"
        )
        assert release_summary_idx < key_highlights_idx
