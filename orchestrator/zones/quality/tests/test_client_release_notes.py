"""Client-facing release notes — PIC/custom only, concise client copy."""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from generators.client_release_notes import (
    PAGE_CONTENT_WIDTH_TWIPS,
    TABLE_INDENT_TWIPS,
    TITLE_SIZE_PT,
    _as_user_control,
    client_notes_filename,
    default_client_notes_path,
    filter_pic_custom_extracted,
    generate_client_release_notes_docx,
    table_left_edges,
    write_client_release_notes_beside,
    write_client_release_notes_file,
)
from generators.release_notes_generator import FONT_NAME
from generators.release_notes_template_defaults import TITLE_COLOR

_FORBIDDEN = (
    "Total issues delivered",
    "issues delivered",
    "What changed for users",
    "Implementation for",
    "User impact:",
    "Expected impact",
    "Validation",
    "Navyasri",
    "Analyze / Standard tickets are not listed here",
    "date date",
    "Commenced at'",
)


def _mixed() -> dict:
    return {
        "version": "v2026.4.10.15",
        "release_date": "15-October-2026",
        "extracted": [
            {
                "ticket": "AN-5596",
                "project": "AN",
                "type": "Analyze",
                "client": "Standard",
                "report": "Class Utilization (dashboard)",
                "description": "Rename Branch column to Location",
            },
            {
                "ticket": "PIC-6326",
                "project": "PIC",
                "type": "Custom",
                "client": "Lift",
                "report": "Corporate Partnership Analysis",
                "description": "Australia update",
                "assignee": "Navyasri Kalla",
                "priority": "Medium",
                "status": "Done",
                "details_plain": (
                    "Tasks -\n"
                    "* Add ‘Discount name’ filter\n"
                    "* Add 'Commenced at' date filter\n"
                ),
            },
            {
                "ticket": "PIC-6058",
                "project": "PIC",
                "type": "Custom",
                "client": "Jetts",
                "report": "Member Transfer (Paginated Report)",
                "description": "additional filters",
                "details_plain": (
                    "Add Membership Commenced From Date and Membership Commenced To Date.\n"
                    "Add Visit Concentration Percentage with options 60%, 70%, and 80%."
                ),
            },
        ],
    }


def _lift_only() -> dict:
    mixed = _mixed()
    mixed["extracted"] = [e for e in mixed["extracted"] if e["ticket"] == "PIC-6326"]
    return mixed


def _doc_text(data: dict) -> str:
    doc = Document(io.BytesIO(generate_client_release_notes_docx(data)))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


def test_filter_drops_analyze():
    kept = filter_pic_custom_extracted(_mixed()["extracted"])
    keys = [e["ticket"] for e in kept]
    assert keys == ["PIC-6326", "PIC-6058"]


def test_repair_country_parked_in_report():
    kept = filter_pic_custom_extracted(
        [
            {
                "ticket": "PIC-6326",
                "project": "PIC",
                "type": "Custom",
                "client": "Lift",
                "report": "Australia",
                "description": "Corporate Partnership Analysis - Update",
            }
        ]
    )
    assert kept[0]["report"] == "Corporate Partnership Analysis"
    assert "Australia" in kept[0]["description"]


def test_client_notes_filename():
    assert client_notes_filename("v2026.4.10.15") == "Client Release Notes 2026.4.10.15.docx"
    assert default_client_notes_path("2026.4.10.15").name == "Client Release Notes 2026.4.10.15.docx"


def test_as_user_control_is_client_language():
    assert _as_user_control("Add ‘Discount name’ filter") == "Filter by Discount name"
    assert _as_user_control("Add 'Commenced at' date filter") == "Choose a commenced date range"
    assert "date date" not in _as_user_control("Add 'Commenced at' date filter")
    assert _as_user_control(
        "Add Visit Concentration Percentage with options 60%, 70%, and 80%."
    ) == "Choose visit concentration (60%, 70%, or 80%)"
    assert _as_user_control(
        "Add Membership Commenced From Date and Membership Commenced To Date."
    ) == "Set a membership commenced date range"


def test_client_notes_sample_headings_and_pic_only():
    text = _doc_text(_mixed())
    doc = Document(io.BytesIO(generate_client_release_notes_docx(_mixed())))
    title_run = doc.paragraphs[0].runs[0]
    assert title_run.text == "Release Notes"
    assert title_run.bold is True
    assert title_run.font.size == Pt(TITLE_SIZE_PT)
    assert "This release includes 2 updates for" in text
    assert "Lift" in text and "Jetts" in text
    assert "Clients in this release" in text
    assert "Discount name" in text
    assert "commenced date range" in text.lower()
    assert "Visit concentration" in text or "visit concentration" in text
    assert "PIC-6326" in text
    assert "PIC-6058" in text
    assert "AN-5596" not in text
    assert "Class Utilization" not in text
    assert "Test by" not in text
    for phrase in _FORBIDDEN:
        assert phrase not in text, phrase
    assert "w:hyperlink" in doc.element.body.xml
    assert doc.styles["Normal"].font.name == FONT_NAME
    assert str(doc.styles["Heading 2"].font.color.rgb) == TITLE_COLOR


def test_single_update_is_short_and_positive():
    text = _doc_text(_lift_only())
    words = [w for w in text.split() if w.strip()]
    assert words
    assert len(words) <= 80
    assert "one update for Lift" in text
    assert "Clients in this release" not in text
    assert "Filter by Discount name" in text
    assert "Choose a commenced date range" in text
    assert "date date" not in text
    for phrase in _FORBIDDEN:
        assert phrase not in text, phrase


def test_tables_share_left_edge_and_width():
    doc = Document(io.BytesIO(generate_client_release_notes_docx(_mixed())))
    # version/date + client overview + Lift banner + Jetts banner
    assert len(doc.tables) == 4
    edges = table_left_edges(doc)
    assert edges == [TABLE_INDENT_TWIPS] * 4
    meta = [[c.text.strip() for c in row.cells] for row in doc.tables[0].rows]
    assert meta[0] == ["Version", "2026.4.10.15"]
    assert meta[1][0] == "Release date"
    overview = [[c.text.strip() for c in row.cells] for row in doc.tables[1].rows]
    assert overview[0] == ["Client", "Reports", "Tickets", "What you get"]
    assert {row[0] for row in overview[1:]} == {"Lift", "Jetts"}
    banner_names = {doc.tables[2].cell(0, 0).text.splitlines()[0], doc.tables[3].cell(0, 0).text.splitlines()[0]}
    assert banner_names == {"Lift", "Jetts"}
    for table in doc.tables:
        tbl_w = table._tbl.tblPr.find(qn("w:tblW"))
        assert tbl_w is not None
        assert int(tbl_w.get(qn("w:w"))) == PAGE_CONTENT_WIDTH_TWIPS


def test_empty_pic_still_builds():
    data = {
        "version": "v2026.4.10.15",
        "release_date": "15-October-2026",
        "extracted": [
            {
                "ticket": "AN-1",
                "type": "Analyze",
                "client": "Standard",
                "report": "Visits",
                "description": "x",
            }
        ],
    }
    text = _doc_text(data)
    assert "no custom-client report updates" in text.lower()
    assert "analyze / standard tickets are not listed here" not in text.lower()
    assert "AN-1" not in text
    assert "Clients in this release" not in text
    assert "Total issues delivered" not in text


def test_write_client_notes_beside_and_file(tmp_path: Path):
    internal = tmp_path / "Release Notes 2026.4.10.15.docx"
    internal.write_bytes(b"internal")
    out = write_client_release_notes_beside(_mixed(), str(internal))
    path = Path(out["client_docx_path"])
    assert path.exists()
    assert out["client_ticket_count"] == 2
    dest = tmp_path / "out" / "client.docx"
    written = write_client_release_notes_file(_mixed(), dest)
    assert Path(written["client_docx_path"]).exists()
