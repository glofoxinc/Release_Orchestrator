"""Built-in release notes DOCX generator using python-docx.

Layout matches the organization reference template (e.g. Release Notes 2026.3.07.15).
"""

from __future__ import annotations

import io
import re
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml.parser import OxmlElement
from docx.shared import Inches, Pt, RGBColor

from generators.parsing import normalize_change_row as _normalize_change_row
from generators.parsing import parse_change_rows as _parse_change_rows
from generators.release_notes_deployment import format_deployment_id_time_docx
from generators.release_notes_docx_styles import patch_docx_grid_table4_accent1_style
from generators.release_notes_report_tables import (
    build_analyze_report_rows,
    build_analyze_report_rows_from_legacy_dashboard_lines,
    build_custom_client_report_tables,
    build_custom_tables_from_legacy_dashboard_lines,
    build_delivered_reports_rows,
    parse_analyze_report_rows,
    parse_custom_client_report_tables,
    normalize_report_row_cells,
)
from generators.release_notes_summary_tables import resolve_highlight_summary_tables
from generators.release_notes_template_defaults import (
    CELL_BORDER_COLOR,
    CELL_BORDER_SZ,
    CHANGE_COLUMN_WIDTHS,
    CHANGE_COLUMN_COUNT,
    CHANGE_DATA_ROW_HEIGHT_TWIPS,
    CHANGE_HEADER_FILL,
    CHANGE_HEADERS,
    CHANGE_PARENT_HEADER_MERGE_WIDTH,
    CHANGE_TICKET_PARENT_MERGE_WIDTH,
    CHANGE_TABLE_WIDTH,
    CHANGE_TBL_IND_TWIPS,
    CHANGE_TBL_LOOK_VAL,
    DELIVERED_REPORTS_COLUMN_WIDTHS,
    DELIVERED_REPORTS_HEADERS,
    DEPLOYMENT_CELL_BORDER_COLOR,
    DEPLOYMENT_CELL_BORDER_SZ,
    DEPLOYMENT_COLUMN_WIDTHS,
    DEPLOYMENT_HEADER_BORDER_COLOR,
    DEPLOYMENT_HEADER_FILL,
    DEPLOYMENT_HEADER_TEXT,
    DEPLOYMENT_HEADERS,
    DEPLOYMENT_TABLE_WIDTH,
    DEPLOYMENT_TBL_IND_TWIPS,
    DEPLOYMENT_TBL_LOOK_VAL,
    EXCLUDED_BODY_FILL,
    EXCLUDED_COLUMN_WIDTHS,
    EXCLUDED_HEADER_FILL,
    EXCLUDED_HEADER_TEXT,
    EXCLUDED_HEADERS,
    EXCLUDED_TABLE_STYLE,
    EXCLUDED_TABLE_WIDTH,
    GRID_TABLE_ACCENT_STYLE,
    HEADING2_SIZE_PT,
    HEADING3_SIZE_PT,
    LIST_STYLE,
    METADATA_GM_LABEL,
    META_LABEL_SUFFIXES,
    PAGE_MARGIN_INCHES,
    REPORT_TABLE_COLUMN_WIDTHS,
    REPORT_TABLE_HEADERS,
    REPORT_TABLE_WIDTH,
    GLOFOX_ANALYZE_CLIENT,
    SECTION_SHADING,
    SCREENSHOTS_DEPLOYMENT_HEADING,
    TESTING_EVIDENCE_HEADING,
    TESTING_EVIDENCE_SUBHEADING,
    SPACING_PROFILES,
    SUMMARY_TABLE_4COL_WIDTH,
    SUMMARY_TABLE_4COL_WIDTHS,
    THEME_ACCENT1_COLOR,
    TITLE_COLOR,
    VERSION_TEXT_COLOR,
)

FONT_NAME = "Cambria"
BODY_SIZE = Pt(11)

GENERATOR_BUILD_ID = "2026.09.09-rn-format-v6-layout"


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def _set_paragraph_shading(paragraph, fill_hex: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:shd"))
    if existing is not None:
        p_pr.remove(existing)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    p_pr.append(shd)


def _set_cell_shading(cell, color_hex: str) -> None:
    tc_pr = cell._element.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:shd"))
    if existing is not None:
        tc_pr.remove(existing)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tc_pr.append(shd)


def _add_cell_borders(cell, *, color: str = CELL_BORDER_COLOR, size: str = CELL_BORDER_SZ) -> None:
    """Per-cell borders matching reference change / excluded tables."""
    tc_pr = cell._element.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:tcBorders"))
    if existing is not None:
        tc_pr.remove(existing)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)
        borders.append(element)
    tc_pr.append(borders)


def _add_deployment_cell_borders(cell, *, header: bool = False) -> None:
    color = DEPLOYMENT_HEADER_BORDER_COLOR if header else DEPLOYMENT_CELL_BORDER_COLOR
    _add_cell_borders(cell, color=color, size=DEPLOYMENT_CELL_BORDER_SZ)


def _set_reference_tbl_look(table, *, look_val: str | None = None) -> None:
    """tblLook flags from Release Notes 2026.3.07.15.docx tables."""
    val = look_val or DEPLOYMENT_TBL_LOOK_VAL
    tbl = table._element
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    existing = tbl_pr.find(qn("w:tblLook"))
    if existing is not None:
        tbl_pr.remove(existing)
    tbl_look = OxmlElement("w:tblLook")
    tbl_look.set(qn("w:val"), val)
    tbl_look.set(qn("w:firstRow"), "1")
    tbl_look.set(qn("w:lastRow"), "0")
    tbl_look.set(qn("w:firstColumn"), "1")
    tbl_look.set(qn("w:lastColumn"), "0")
    tbl_look.set(qn("w:noHBand"), "1")
    tbl_look.set(qn("w:noVBand"), "1")
    tbl_pr.append(tbl_look)


def _set_deployment_header_row_cnf(row) -> None:
    _set_table_row_cnf(row, first_row=True)
    _set_table_row_header_repeat(row, enabled=True)


def _set_change_header_row_cnf(row) -> None:
    _set_table_row_cnf(row, first_row=True)
    _set_table_row_header_repeat(row, enabled=True)


def _set_table_row_cnf(row, *, first_row: bool) -> None:
    tr = row._tr
    tr_pr = tr.get_or_add_trPr()
    existing = tr_pr.find(qn("w:cnfStyle"))
    if existing is not None:
        tr_pr.remove(existing)
    cnf = OxmlElement("w:cnfStyle")
    cnf.set(qn("w:val"), "100000000000" if first_row else "000000000000")
    cnf.set(qn("w:firstRow"), "1" if first_row else "0")
    cnf.set(qn("w:lastRow"), "0")
    cnf.set(qn("w:firstColumn"), "0")
    cnf.set(qn("w:lastColumn"), "0")
    cnf.set(qn("w:oddVBand"), "0")
    cnf.set(qn("w:evenVBand"), "0")
    cnf.set(qn("w:oddHBand"), "0")
    cnf.set(qn("w:evenHBand"), "0")
    cnf.set(qn("w:firstRowFirstColumn"), "0")
    cnf.set(qn("w:firstRowLastColumn"), "0")
    cnf.set(qn("w:lastRowFirstColumn"), "0")
    cnf.set(qn("w:lastRowLastColumn"), "0")
    tr_pr.append(cnf)


def _set_table_row_header_repeat(row, *, enabled: bool = True) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    existing = tr_pr.find(qn("w:tblHeader"))
    if enabled:
        if existing is None:
            tr_pr.append(OxmlElement("w:tblHeader"))
    elif existing is not None:
        tr_pr.remove(existing)


def _set_cell_paragraph_cnf(cell, *, first_row: bool) -> None:
    paragraph = cell.paragraphs[0]
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:cnfStyle"))
    if existing is not None:
        p_pr.remove(existing)
    cnf = OxmlElement("w:cnfStyle")
    cnf.set(qn("w:val"), "100000000000" if first_row else "000000000000")
    cnf.set(qn("w:firstRow"), "1" if first_row else "0")
    cnf.set(qn("w:lastRow"), "0")
    cnf.set(qn("w:firstColumn"), "0")
    cnf.set(qn("w:lastColumn"), "0")
    cnf.set(qn("w:oddVBand"), "0")
    cnf.set(qn("w:evenVBand"), "0")
    cnf.set(qn("w:oddHBand"), "0")
    cnf.set(qn("w:evenHBand"), "0")
    cnf.set(qn("w:firstRowFirstColumn"), "0")
    cnf.set(qn("w:firstRowLastColumn"), "0")
    cnf.set(qn("w:lastRowFirstColumn"), "0")
    cnf.set(qn("w:lastRowLastColumn"), "0")
    p_pr.append(cnf)


def _set_cell_paragraph_font_rpr(cell) -> None:
    """Reference table cells carry Cambria 11 (w:sz 22) in paragraph rPr."""
    paragraph = cell.paragraphs[0]
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:rPr"))
    if existing is not None:
        p_pr.remove(existing)
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT_NAME)
    fonts.set(qn("w:hAnsi"), FONT_NAME)
    r_pr.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "22")
    r_pr.append(size)
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), "22")
    r_pr.append(size_cs)
    p_pr.append(r_pr)


def _set_change_data_row_height(row) -> None:
    tr = row._tr
    tr_pr = tr.get_or_add_trPr()
    existing = tr_pr.find(qn("w:trHeight"))
    if existing is not None:
        tr_pr.remove(existing)
    tr_height = OxmlElement("w:trHeight")
    tr_height.set(qn("w:val"), str(CHANGE_DATA_ROW_HEIGHT_TWIPS))
    tr_pr.append(tr_height)


def _add_reference_gap_paragraphs(doc: Document, *, count_240: int = 2, count_278: int = 1) -> None:
    """Empty paragraphs before deployment / after change table (reference spacing)."""
    for _ in range(count_240):
        gap = doc.add_paragraph("")
        _format_section_paragraph(gap, spacing="dashboard_empty_gap")
    for _ in range(count_278):
        gap = doc.add_paragraph("")
        _format_section_paragraph(gap, spacing="pre_deployment_gap")


def _add_post_deployment_table_gaps(doc: Document) -> None:
    """Breathing room between deployment table and Change Details section."""
    gap = doc.add_paragraph("")
    _format_section_paragraph(gap, spacing="post_deployment_table_gap")


def _add_pre_change_table_gaps(doc: Document) -> None:
    """Breathing room between Change Details heading and change table."""
    gap = doc.add_paragraph("")
    _format_section_paragraph(gap, spacing="pre_change_table_gap")


def _add_post_change_table_gaps(doc: Document) -> None:
    """Spacing block between change table and sanity checklist in reference."""
    doc.add_paragraph("")
    doc.add_paragraph("")
    nbsp_gap = doc.add_paragraph("")
    nbsp_gap.add_run("\u00a0")
    _format_section_paragraph(nbsp_gap, spacing="post_change_table_gap")
    doc.add_paragraph("")
    post_change_gap = doc.add_paragraph("")
    _format_section_paragraph(post_change_gap, spacing="post_change_table_gap")


def _set_table_indent(table, twips: int) -> None:
    tbl = table._element
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    existing = tbl_pr.find(qn("w:tblInd"))
    if existing is not None:
        tbl_pr.remove(existing)
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), str(twips))
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)


def _set_deployment_table_indent(table) -> None:
    _set_table_indent(table, DEPLOYMENT_TBL_IND_TWIPS)


def _set_change_table_indent(table) -> None:
    _set_table_indent(table, CHANGE_TBL_IND_TWIPS)


def _set_cell_tc_width(cell, width: int) -> None:
    tc_pr = cell._element.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:tcW"))
    if existing is not None:
        tc_pr.remove(existing)
    tc_w = OxmlElement("w:tcW")
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")
    tc_pr.append(tc_w)


def _apply_change_table_merges(table) -> None:
    """Legacy 10-column Parent merges — only for the old wide change table layout."""
    if CHANGE_COLUMN_COUNT != 10:
        return
    header = table.rows[0]
    parent_left = header.cells[2]
    parent_right = header.cells[3]
    parent_left.merge(parent_right)
    parent_left.text = "Parent"
    for paragraph in parent_left.paragraphs:
        for run in paragraph.runs:
            _style_run(run, size=BODY_SIZE)
    _set_cell_shading(parent_left, CHANGE_HEADER_FILL)
    _add_cell_borders(parent_left)
    _set_cell_tc_width(parent_left, CHANGE_PARENT_HEADER_MERGE_WIDTH)

    for row in table.rows[1:]:
        ticket_cell = row.cells[1]
        parent_ticket_cell = row.cells[2]
        ticket_value = ticket_cell.text
        ticket_cell.merge(parent_ticket_cell)
        if ticket_cell.text.strip() != ticket_value.strip():
            ticket_cell.text = ticket_value
        _set_cell_tc_width(ticket_cell, CHANGE_TICKET_PARENT_MERGE_WIDTH)
        _add_cell_borders(ticket_cell)


def _style_section_heading_run(run, *, level: int = 2) -> None:
    """Heading 2/3 labels: org blue #0F4761, larger size — do not flatten to body Cambria 11."""
    run.bold = True
    run.font.color.rgb = _rgb(TITLE_COLOR)
    run.font.size = Pt(HEADING2_SIZE_PT if level == 2 else HEADING3_SIZE_PT)


def _add_shaded_section_heading(doc: Document, text: str, *, level: int = 2):
    para = doc.add_paragraph(style=f"Heading {level}")
    run = para.add_run(text)
    _style_section_heading_run(run, level=level)
    spacing = "shaded_section" if level == 2 else "shaded_heading3"
    _format_section_paragraph(para, shaded=True, spacing=spacing)
    return para


def _split_dashboard_line(line: str) -> tuple[str, str]:
    """Split dashboard line into bold client prefix and remainder."""
    if " - " in line:
        client, rest = line.split(" - ", 1)
        return client.strip(), f" - {rest}"
    dash = line.find("-")
    if dash > 0:
        return line[:dash].strip(), line[dash:]
    return line, ""


def _write_dashboard_line(paragraph, line: str) -> None:
    prefix, suffix = _split_dashboard_line(line)
    if prefix and suffix:
        lead = paragraph.add_run(prefix)
        _style_run(lead, bold=True)
        tail = paragraph.add_run(suffix)
        _style_run(tail)
    else:
        run = paragraph.add_run(line)
        _style_run(run)


def _style_run(
    run,
    *,
    bold: bool | None = None,
    italic: bool | None = None,
    size=None,
    color_hex: str | None = None,
) -> None:
    run.font.name = FONT_NAME
    run.font.size = size or BODY_SIZE
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color_hex:
        run.font.color.rgb = _rgb(color_hex)


def _shade_paragraph(paragraph) -> None:
    _set_paragraph_shading(paragraph, SECTION_SHADING)


def _set_paragraph_spacing(
    paragraph,
    *,
    before: str | None = None,
    after: str | None = None,
    line: str | None = None,
    line_rule: str = "auto",
) -> None:
    """Apply w:spacing values matching the reference release-notes template."""
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:spacing"))
    if existing is not None:
        p_pr.remove(existing)
    if before is None and after is None and line is None:
        return
    spacing = OxmlElement("w:spacing")
    if before is not None:
        spacing.set(qn("w:before"), before)
    if after is not None:
        spacing.set(qn("w:after"), after)
    if line is not None:
        spacing.set(qn("w:line"), line)
        spacing.set(qn("w:lineRule"), line_rule)
    p_pr.append(spacing)


def _apply_spacing_profile(paragraph, profile: str) -> None:
    spec = dict(SPACING_PROFILES[profile])
    line_rule = spec.pop("lineRule", "auto")
    _set_paragraph_spacing(paragraph, line_rule=line_rule, **spec)


def _format_section_paragraph(
    paragraph,
    *,
    shaded: bool = False,
    spacing: str | None = None,
) -> None:
    if shaded:
        _shade_paragraph(paragraph)
    if spacing:
        _apply_spacing_profile(paragraph, spacing)


def _set_paragraph_keep_with_next(paragraph, *, enabled: bool = True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:keepNext"))
    if enabled:
        if existing is None:
            p_pr.append(OxmlElement("w:keepNext"))
    elif existing is not None:
        p_pr.remove(existing)


def _set_paragraph_keep_lines(paragraph, *, enabled: bool = True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:keepLines"))
    if enabled:
        if existing is None:
            p_pr.append(OxmlElement("w:keepLines"))
    elif existing is not None:
        p_pr.remove(existing)


def _set_paragraph_page_break_before(paragraph, *, enabled: bool = True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:pageBreakBefore"))
    if enabled:
        if existing is None:
            p_pr.append(OxmlElement("w:pageBreakBefore"))
    elif existing is not None:
        p_pr.remove(existing)


def _set_table_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def _prevent_table_row_splits(table) -> None:
    for row in table.rows:
        _set_table_row_cant_split(row)


def _add_doc_spacer(doc: Document, spacing: str) -> None:
    spacer = doc.add_paragraph("")
    _format_section_paragraph(spacer, spacing=spacing)


def _bind_block(paragraph, *companions) -> None:
    """Keep a heading or label with the content that follows it."""
    _set_paragraph_keep_with_next(paragraph)
    for companion in companions:
        _set_paragraph_keep_lines(companion)


def _configure_reference_heading_styles(doc: Document) -> None:
    """Override python-docx default heading blue (#4F81BD) with org color #0F4761."""
    doc.styles["Heading 2"].font.color.rgb = _rgb(TITLE_COLOR)
    doc.styles["Heading 2"].font.size = Pt(HEADING2_SIZE_PT)
    doc.styles["Heading 2"].font.bold = True
    doc.styles["Heading 3"].font.color.rgb = _rgb(TITLE_COLOR)
    doc.styles["Heading 3"].font.size = Pt(HEADING3_SIZE_PT)
    doc.styles["Heading 3"].font.bold = True
    doc.styles["Heading 1"].font.name = FONT_NAME
    doc.styles["Heading 1"].font.color.rgb = _rgb(TITLE_COLOR)


def _patch_document_theme_accent1(doc: Document, accent_hex: str) -> None:
    """Match reference table-style tint (Grid Table 4 Accent 1 uses theme accent1)."""
    theme_reltype = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme"
    )
    try:
        theme_part = doc.part.part_related_by(theme_reltype)
    except KeyError:
        return
    xml = theme_part.blob.decode("utf-8")
    patched = re.sub(
        r"(<a:accent1>.*?<a:srgbClr val=\")[^\"]+(\")",
        rf"\g<1>{accent_hex.upper()}\2",
        xml,
        count=1,
        flags=re.DOTALL,
    )
    if patched != xml:
        theme_part._blob = patched.encode("utf-8")


def _configure_reference_document_theme(doc: Document) -> None:
    _configure_reference_heading_styles(doc)
    _patch_document_theme_accent1(doc, THEME_ACCENT1_COLOR)


def _add_table_borders(table) -> None:
    """Reference excluded-tickets placeholder uses table-level borders."""
    tbl = table._element
    tbl_pr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "BFBFBF")
        borders.append(element)
    tbl_pr.append(borders)
    if tbl.tblPr is None:
        tbl.insert(0, tbl_pr)


def _set_cell_no_wrap(cell) -> None:
    tc_pr = cell._element.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:noWrap"))
    if existing is not None:
        tc_pr.remove(existing)
    no_wrap = OxmlElement("w:noWrap")
    tc_pr.append(no_wrap)


def _apply_table_column_widths(
    table,
    widths: tuple[int, ...],
    *,
    total_width: int | None = None,
    fixed_layout: bool = False,
) -> None:
    """Apply column widths (twips/dxa). Change table uses fixed layout; others match reference."""
    try:
        table.autofit = False
    except AttributeError:
        pass

    tbl = table._element
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)

    layout = tbl_pr.find(qn("w:tblLayout"))
    if fixed_layout:
        if layout is None:
            layout = OxmlElement("w:tblLayout")
            tbl_pr.append(layout)
        layout.set(qn("w:type"), "fixed")
    elif layout is not None:
        tbl_pr.remove(layout)

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if total_width is not None:
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW")
            tbl_pr.append(tbl_w)
        tbl_w.set(qn("w:w"), str(total_width))
        tbl_w.set(qn("w:type"), "dxa")
    elif tbl_w is not None:
        tbl_w.set(qn("w:w"), "0")
        tbl_w.set(qn("w:type"), "auto")
    else:
        tbl_w = OxmlElement("w:tblW")
        tbl_w.set(qn("w:w"), "0")
        tbl_w.set(qn("w:type"), "auto")
        tbl_pr.append(tbl_w)

    grid = tbl.find(qn("w:tblGrid"))
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(1, grid)
    for existing in list(grid.findall(qn("w:gridCol"))):
        grid.remove(existing)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            if idx >= len(widths):
                break
            tc_pr = cell._element.get_or_add_tcPr()
            existing = tc_pr.find(qn("w:tcW"))
            if existing is not None:
                tc_pr.remove(existing)
            tc_w = OxmlElement("w:tcW")
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            tc_pr.append(tc_w)


def _split_version_run_parts(version: str) -> list[str]:
    """Split version text like the reference DOCX (per-segment runs)."""
    clean = version.lstrip("vV")
    match = re.match(r"(\d{4})\.(\d)\.(\d{2})\.(\d{2})$", clean)
    if not match:
        prefix = version if version.lower().startswith("v") else f"v{version}"
        return [f" Version: {prefix}"]
    year, quarter, month, day = match.groups()
    return [
        f" Version: v{year}.",
        quarter,
        f".{month[0]}",
        month[1],
        ".",
        day,
    ]


def _style_change_header_row(row) -> None:
    """Reference change table: black header cells, Cambria 11, not bold, inherited text color."""
    for cell in row.cells:
        _set_cell_shading(cell, CHANGE_HEADER_FILL)
        _add_cell_borders(cell)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                _style_run(run, size=BODY_SIZE)


def _apply_reference_margins(doc: Document) -> None:
    for section in doc.sections:
        section.left_margin = Inches(PAGE_MARGIN_INCHES)
        section.right_margin = Inches(PAGE_MARGIN_INCHES)
        section.top_margin = Inches(PAGE_MARGIN_INCHES)
        section.bottom_margin = Inches(PAGE_MARGIN_INCHES)


def _add_label_value_line(doc: Document, label: str, value: str) -> None:
    """Metadata line with reference tab alignment; label, colon, and value are bold."""
    para = doc.add_paragraph()
    label_text = label if label in META_LABEL_SUFFIXES else label.rstrip()
    suffix = META_LABEL_SUFFIXES.get(label, META_LABEL_SUFFIXES.get(label_text, "\t\t: \t"))
    label_run = para.add_run(label_text)
    _style_run(label_run, bold=True, color_hex=VERSION_TEXT_COLOR)
    if suffix.endswith(": \t"):
        prefix, tail = suffix.rsplit(": \t", 1)
        sep_prefix = para.add_run(prefix)
        _style_run(sep_prefix, bold=False)
        sep_colon = para.add_run(": \t")
        _style_run(sep_colon, bold=True, color_hex=VERSION_TEXT_COLOR)
    else:
        sep_run = para.add_run(suffix)
        _style_run(sep_run, bold=False)
    value_run = para.add_run(value)
    _style_run(value_run, bold=True, color_hex=VERSION_TEXT_COLOR)


def _parse_lines(raw: str) -> list[str]:
    return [line.strip() for line in (raw or "").strip().splitlines() if line.strip()]


def _parse_key_highlights(raw: str) -> list[tuple[str, str]]:
    items: list[tuple[str, str]] = []
    for line in _parse_lines(raw):
        if ":" in line:
            prefix, rest = line.split(":", 1)
            items.append((prefix.strip(), rest.strip()))
        else:
            items.append(("", line))
    return items


def build_key_highlight_lines(extracted: list[dict[str, Any]]) -> list[str]:
    """Build polished 'Key Highlights' lines from extracted ticket data.

    Delegates to :mod:`generators.release_notes_quality` so typo cleanup,
    weak-report repair, and client grouping stay in one place.
    """
    from generators.release_notes_quality import (
        build_key_highlight_lines as _build_key_highlight_lines,
    )

    return _build_key_highlight_lines(extracted)


def _parse_dashboard_lines(raw: str) -> list[str]:
    lines = _parse_lines(raw)
    if not lines:
        return []
    items: list[str] = []
    for line in lines:
        if " - " in line and not line.lower().startswith("http"):
            parts = line.split(" - ", 1)
            items.append(f"{parts[0].strip()} - {parts[1].strip()}")
        else:
            items.append(line)
    return items


def _parse_pipe_rows(raw: str, *, min_cols: int) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in _parse_lines(raw):
        cols = [c.strip() for c in line.split("|", min_cols - 1)]
        while len(cols) < min_cols:
            cols.append("")
        rows.append(cols[:min_cols])
    return rows


def format_change_row(
    *,
    sno: str = "",
    ticket: str,
    parent_epic: str = "",
    link: str,
    report: str = "",
    description: str = "",
    test_by: str,
    assignee: str = "",
    comments: str = "",
    testing_status: str = "",
    type_: str = "",
    client: str = "",
) -> str:
    """Build a 5-column pipe row for the Ticket Details table."""
    del sno, parent_epic, report, description, comments, type_, client
    resolved_link = (link or "").strip()
    if not resolved_link.startswith(("http://", "https://")) and ticket and "-" in ticket:
        try:
            import config

            base = (config.JIRA_BASE_URL or "").rstrip("/")
            if base:
                resolved_link = f"{base}/browse/{ticket}"
        except (ImportError, AttributeError):
            pass
    return "|".join([ticket, resolved_link, assignee, test_by, testing_status])


_TABLE_STYLE_IDS = {
    GRID_TABLE_ACCENT_STYLE: "GridTable4-Accent1",
    EXCLUDED_TABLE_STYLE: "TableGrid",
}


def _apply_reference_table_style(table, style_name: str) -> None:
    """Apply the Word table style used in the org reference template."""
    try:
        table.style = style_name
    except KeyError:
        pass
    tbl = table._element
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    existing = tbl_pr.find(qn("w:tblStyle"))
    if existing is not None:
        tbl_pr.remove(existing)
    style_el = OxmlElement("w:tblStyle")
    style_el.set(qn("w:val"), _TABLE_STYLE_IDS.get(style_name, style_name))
    tbl_pr.append(style_el)


def _reset_cell_paragraph(cell) -> None:
    """Remove all runs from the first paragraph in a table cell."""
    _clear_cell_paragraph_content(cell)


def _clear_cell_paragraph_content(cell) -> None:
    """Clear every child of the first paragraph except ``w:pPr`` (runs and hyperlinks)."""
    cell.text = ""
    paragraph = cell.paragraphs[0]
    p = paragraph._p
    p_pr = p.find(qn("w:pPr"))
    for child in list(p):
        if child is not p_pr:
            p.remove(child)


def _write_empty_cell(cell) -> None:
    """Empty table cell with no runs (matches reference Status / ID placeholders)."""
    _reset_cell_paragraph(cell)


def _write_deployment_header_cell(cell, header: str) -> None:
    """Reference deployment header: blue fill, white centered Cambria 11."""
    _reset_cell_paragraph(cell)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if header == "Deployment Owner":
        part1 = paragraph.add_run("Deployment ")
        _style_run(part1, size=BODY_SIZE, color_hex=DEPLOYMENT_HEADER_TEXT)
        part2 = paragraph.add_run("Owner")
        _style_run(part2, size=BODY_SIZE, color_hex=DEPLOYMENT_HEADER_TEXT)
    else:
        run = paragraph.add_run(header)
        _style_run(run, size=BODY_SIZE, color_hex=DEPLOYMENT_HEADER_TEXT)
    _set_cell_shading(cell, DEPLOYMENT_HEADER_FILL)
    _add_deployment_cell_borders(cell, header=True)


def _write_deployment_data_cell(cell, value: str, *, column: int) -> None:
    """Reference deployment body cells: client left; owner centered bold; status/id centered."""
    text = (value or "").strip()
    if not text:
        _write_empty_cell(cell)
        _add_deployment_cell_borders(cell, header=False)
        return

    _reset_cell_paragraph(cell)
    paragraph = cell.paragraphs[0]
    if column == 0:
        run = paragraph.add_run(text)
        _style_run(run, size=BODY_SIZE)
    elif column == 1:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(f" {text}")
        _style_run(run, size=BODY_SIZE, bold=True)
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(value)
        _style_run(run, size=BODY_SIZE)
    _add_deployment_cell_borders(cell, header=False)


def _write_deployment_table(table, deployment_rows: list[list[str]]) -> None:
    _apply_reference_table_style(table, GRID_TABLE_ACCENT_STYLE)
    _set_reference_tbl_look(table)
    _set_deployment_table_indent(table)
    header_row = table.rows[0]
    _set_deployment_header_row_cnf(header_row)
    for col_idx, header in enumerate(DEPLOYMENT_HEADERS):
        _write_deployment_header_cell(header_row.cells[col_idx], header)

    for row_idx, row_data in enumerate(deployment_rows):
        row = table.rows[row_idx + 1]
        for col_idx, value in enumerate(row_data):
            display = format_deployment_id_time_docx(value) if col_idx == 3 else value
            _write_deployment_data_cell(row.cells[col_idx], display, column=col_idx)

    _apply_table_column_widths(
        table,
        DEPLOYMENT_COLUMN_WIDTHS,
        total_width=DEPLOYMENT_TABLE_WIDTH,
        fixed_layout=False,
    )
    _prevent_table_row_splits(table)


def _detail_rows_from_inventory(rows: list[list[str]]) -> list[list[str]]:
    """Drop Client column — used under per-client headings."""
    detail: list[list[str]] = []
    for row in rows:
        if len(row) >= 6:
            detail.append(row[1:6])
        elif len(row) >= 5:
            detail.append(row[:5])
        else:
            detail.append(row)
    return detail


def _write_report_inventory_table(
    table,
    rows: list[list[str]],
    *,
    headers: tuple[str, ...] | list[str] | None = None,
    column_widths: tuple[int, ...] | None = None,
    indent_twips: int | None = None,
) -> None:
    """Dashboard / paginated inventory table."""
    header_list = list(headers or REPORT_TABLE_HEADERS)
    widths = column_widths or REPORT_TABLE_COLUMN_WIDTHS
    table_width = sum(widths) if widths else REPORT_TABLE_WIDTH
    _apply_reference_table_style(table, GRID_TABLE_ACCENT_STYLE)
    _set_reference_tbl_look(table)
    if indent_twips is not None:
        _set_table_indent(table, indent_twips)
    header_row = table.rows[0]
    _set_deployment_header_row_cnf(header_row)
    for col_idx, header in enumerate(header_list):
        _write_deployment_header_cell(header_row.cells[col_idx], header)

    body_rows = rows or [[]]
    for row_idx, row_data in enumerate(body_rows):
        row = table.rows[row_idx + 1]
        for col_idx in range(len(header_list)):
            value = row_data[col_idx] if col_idx < len(row_data) else ""
            _write_deployment_data_cell(row.cells[col_idx], value, column=col_idx)

    _apply_table_column_widths(
        table,
        widths,
        total_width=table_width,
        fixed_layout=False,
    )
    _prevent_table_row_splits(table)


def _add_report_inventory_table(
    doc,
    rows: list[list[str]],
    *,
    headers: tuple[str, ...] | list[str] | None = None,
    column_widths: tuple[int, ...] | None = None,
    indent_twips: int | None = None,
) -> None:
    header_list = list(headers or REPORT_TABLE_HEADERS)
    table = doc.add_table(rows=1 + max(len(rows), 1), cols=len(header_list))
    _write_report_inventory_table(
        table,
        rows,
        headers=header_list,
        column_widths=column_widths,
        indent_twips=indent_twips,
    )


def _resolve_analyze_report_rows(data: dict[str, Any]) -> list[list[str]]:
    rows = parse_analyze_report_rows(data.get("analyze_report_rows", ""))
    if rows:
        return rows
    legacy = _parse_dashboard_lines(data.get("dashboard_items", ""))
    if legacy:
        return [row.to_row() for row in build_analyze_report_rows_from_legacy_dashboard_lines(legacy)]
    extracted = data.get("extracted") or []
    if extracted:
        return [row.to_row() for row in build_analyze_report_rows(extracted)]
    return []


def _resolve_custom_client_report_tables(data: dict[str, Any]) -> dict[str, list[list[str]]]:
    tables = parse_custom_client_report_tables(data.get("custom_report_tables", ""))
    if tables:
        return tables
    legacy = _parse_dashboard_lines(data.get("dashboard_items", ""))
    if legacy:
        built = build_custom_tables_from_legacy_dashboard_lines(legacy)
        return {client: [row.to_row() for row in rows] for client, rows in built.items()}
    extracted = data.get("extracted") or []
    if extracted:
        built = build_custom_client_report_tables(extracted)
        return {client: [row.to_row() for row in rows] for client, rows in built.items()}
    custom_clients = [
        c.strip() for c in str(data.get("custom_clients") or "").split(",") if c.strip()
    ]
    return {client: [] for client in custom_clients}


def _resolve_delivered_reports_rows(data: dict[str, Any]) -> list[list[str]]:
    raw = data.get("delivered_report_rows", "")
    if raw:
        return _parse_pipe_rows(raw, min_cols=3)

    extracted = data.get("extracted") or []
    if extracted:
        return build_delivered_reports_rows(extracted)

    rows: list[list[str]] = []
    for inv in _resolve_analyze_report_rows(data):
        normalized = normalize_report_row_cells(inv, default_client="Analyze")
        change = normalized[5] if len(normalized) > 5 else ""
        if normalized[1]:
            rows.append([GLOFOX_ANALYZE_CLIENT, normalized[1], change])

    for client, inv_rows in _resolve_custom_client_report_tables(data).items():
        for inv in inv_rows:
            normalized = normalize_report_row_cells(inv, default_client=client)
            change = normalized[5] if len(normalized) > 5 else ""
            if normalized[1]:
                rows.append([client, normalized[1], change])
    return rows


def _write_summary_table(table, headers: tuple[str, ...] | list[str], rows: list[list[str]]) -> None:
    """Generic styled table for Key Highlights summary blocks."""
    _apply_reference_table_style(table, GRID_TABLE_ACCENT_STYLE)
    _set_reference_tbl_look(table)
    header_row = table.rows[0]
    _set_deployment_header_row_cnf(header_row)
    header_list = list(headers)
    for col_idx, header in enumerate(header_list):
        _write_deployment_header_cell(header_row.cells[col_idx], header)

    body_rows = rows or [[]]
    for row_idx, row_data in enumerate(body_rows):
        row = table.rows[row_idx + 1]
        for col_idx in range(len(header_list)):
            value = row_data[col_idx] if col_idx < len(row_data) else ""
            _write_deployment_data_cell(row.cells[col_idx], value, column=col_idx)

    widths = SUMMARY_TABLE_4COL_WIDTHS
    if len(header_list) != 4:
        per_col = max(SUMMARY_TABLE_4COL_WIDTH // max(len(header_list), 1), 300)
        widths = tuple(per_col for _ in header_list)
        total = sum(widths)
    else:
        total = SUMMARY_TABLE_4COL_WIDTH
    _apply_table_column_widths(table, widths, total_width=total, fixed_layout=False)
    _prevent_table_row_splits(table)


def _add_summary_table(doc, headers: tuple[str, ...] | list[str], rows: list[list[str]]) -> None:
    header_list = list(headers)
    table = doc.add_table(rows=1 + max(len(rows), 1), cols=len(header_list))
    _write_summary_table(table, header_list, rows)


def _write_table_cell(
    cell,
    value: str,
    *,
    bordered: bool = False,
    color_hex: str | None = VERSION_TEXT_COLOR,
) -> None:
    cell.text = value
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            if color_hex:
                _style_run(run, size=BODY_SIZE, color_hex=color_hex)
            else:
                _style_run(run, size=BODY_SIZE)
    if bordered:
        _add_cell_borders(cell)


HYPERLINK_COLOR = "0563C1"
_HYPERLINK_RELATIONSHIP = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"


def _derive_ticket_url(reference_url: str, ticket_key: str) -> str:
    """Build a sibling ticket URL from another ticket's browse link.

    ``https://org.atlassian.net/browse/PIC-5196`` + ``PIC-4000`` ->
    ``https://org.atlassian.net/browse/PIC-4000``. Falls back to the
    reference URL when either input is missing.
    """
    reference_url = (reference_url or "").strip()
    ticket_key = (ticket_key or "").strip()
    if not reference_url or not ticket_key:
        return reference_url
    base = reference_url.rsplit("/", 1)[0]
    return f"{base}/{ticket_key}"


def _add_hyperlink_run(paragraph, text: str, url: str) -> None:
    """Insert a clickable `w:hyperlink` run (python-docx has no writer for this)."""
    part = paragraph.part
    r_id = part.relate_to(url, _HYPERLINK_RELATIONSHIP, is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run_el = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT_NAME)
    fonts.set(qn("w:hAnsi"), FONT_NAME)
    r_pr.append(fonts)
    for tag in ("sz", "szCs"):
        size_el = OxmlElement(f"w:{tag}")
        size_el.set(qn("w:val"), "22")
        r_pr.append(size_el)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), HYPERLINK_COLOR)
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    run_el.append(r_pr)

    t = OxmlElement("w:t")
    if text.startswith(" ") or text.endswith(" "):
        t.set(qn("xml:space"), "preserve")
    t.text = text
    run_el.append(t)
    hyperlink.append(run_el)
    paragraph._p.append(hyperlink)


def _write_table_cell_link(
    cell,
    text: str,
    url: str,
    *,
    bordered: bool = False,
) -> None:
    """Write a cell whose visible text is a clickable hyperlink to ``url``."""
    text = (text or "").strip()
    url = (url or "").strip()
    _clear_cell_paragraph_content(cell)
    paragraph = cell.paragraphs[0]
    if text and url:
        _add_hyperlink_run(paragraph, text, url)
    elif text:
        run = paragraph.add_run(text)
        _style_run(run, size=BODY_SIZE)
    if bordered:
        _add_cell_borders(cell)


def resolve_change_row_link_url(row_data: list[str]) -> str:
    """Return the Jira browse URL for a ticket-details row."""
    if len(row_data) >= 4 and len(row_data) <= CHANGE_COLUMN_COUNT:
        explicit = (row_data[1] or "").strip()
        if explicit.startswith(("http://", "https://")):
            return explicit
        ticket = (row_data[0] or "").strip()
    else:
        explicit = (row_data[4] or "").strip() if len(row_data) > 4 else ""
        if explicit.startswith(("http://", "https://")):
            return explicit
        ticket = (row_data[1] or "").strip() if len(row_data) > 1 else ""

    for value in row_data:
        candidate = (value or "").strip()
        if candidate.startswith(("http://", "https://")):
            return candidate

    if ticket and "-" in ticket:
        try:
            import config

            base = (config.JIRA_BASE_URL or "").rstrip("/")
            if base:
                return f"{base}/browse/{ticket}"
        except (ImportError, AttributeError):
            pass
    return ""


def count_change_table_hyperlinks(docx_bytes: bytes) -> int:
    """Count ``w:hyperlink`` elements in Ticket and Ticket Link columns."""
    doc = Document(io.BytesIO(docx_bytes))
    table = next(
        (
            t
            for t in doc.tables
            if any("Ticket" in c.text for c in t.rows[0].cells)
            and len(t.columns) == CHANGE_COLUMN_COUNT
        ),
        None,
    )
    if table is None:
        return 0
    total = 0
    for row in table.rows[1:]:
        for col_idx in (0, 1):
            if col_idx < len(row.cells):
                total += len(row.cells[col_idx].paragraphs[0]._p.findall(qn("w:hyperlink")))
    return total


def _write_change_table(table, change_rows: list[list[str]]) -> None:
    _apply_reference_table_style(table, GRID_TABLE_ACCENT_STYLE)
    _set_reference_tbl_look(table, look_val=CHANGE_TBL_LOOK_VAL)
    _set_change_table_indent(table)
    header_row = table.rows[0]
    _set_change_header_row_cnf(header_row)
    for i, header_text in enumerate(CHANGE_HEADERS):
        cell = header_row.cells[i]
        cell.text = header_text
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                _style_run(run, size=BODY_SIZE)
        _set_cell_no_wrap(cell)
        _set_cell_paragraph_cnf(cell, first_row=True)
        _set_cell_paragraph_font_rpr(cell)
        _add_cell_borders(cell)
    _style_change_header_row(header_row)

    for row_idx, row_data in enumerate(change_rows):
        row = table.rows[row_idx + 1]
        _set_change_data_row_height(row)
        for col_idx, value in enumerate(row_data):
            _write_table_cell(row.cells[col_idx], value, bordered=True, color_hex=None)
        for cell in row.cells:
            _set_cell_paragraph_cnf(cell, first_row=False)
            _set_cell_paragraph_font_rpr(cell)

    _apply_table_column_widths(
        table, CHANGE_COLUMN_WIDTHS, total_width=CHANGE_TABLE_WIDTH, fixed_layout=True
    )
    _apply_change_table_merges(table)
    if CHANGE_COLUMN_COUNT > 4:
        _set_cell_paragraph_cnf(header_row.cells[2], first_row=True)
        _set_cell_paragraph_font_rpr(header_row.cells[2])
        _set_cell_no_wrap(header_row.cells[2])
        for row in table.rows[1:]:
            _set_cell_paragraph_cnf(row.cells[1], first_row=False)
            _set_cell_paragraph_font_rpr(row.cells[1])

    _apply_change_table_hyperlinks(table, change_rows)
    _prevent_table_row_splits(table)


def _apply_change_table_hyperlinks(table, change_rows: list[list[str]]) -> None:
    """Turn Ticket and Ticket Link cells into clickable Jira hyperlinks."""
    for row, row_data in zip(table.rows[1:], change_rows):
        link_url = resolve_change_row_link_url(row_data)
        if not link_url:
            continue

        ticket_value = (row_data[0] or "").strip() if len(row_data) > 0 else ""
        if ticket_value:
            _write_table_cell_link(row.cells[0], ticket_value, link_url, bordered=True)
            _set_cell_paragraph_cnf(row.cells[0], first_row=False)

        link_label = (row_data[1] or "").strip() if len(row_data) > 1 else ""
        if not link_label:
            link_label = link_url
        _write_table_cell_link(row.cells[1], link_label, link_url, bordered=True)
        _set_cell_paragraph_cnf(row.cells[1], first_row=False)


def _testing_evidence_ticket_lines(data: dict[str, Any]) -> list[str]:
    """Ticket keys listed under Testing Evidence (testers paste images below each line)."""
    raw = (
        data.get("testing_evidence_tickets", "")
        or data.get("screenshot_tickets", "")
        or data.get("screenshot_names", "")
    )
    lines = _parse_lines(raw)
    if lines:
        in_ticket_section = False
        tickets: list[str] = []
        for line in lines:
            stripped = line.strip()
            if not stripped:
                continue
            if stripped.lower() == "tickets list":
                in_ticket_section = True
                continue
            if stripped.startswith("Add ") or stripped.startswith("  Add "):
                continue
            if not in_ticket_section and "|" not in stripped and "-" not in stripped:
                continue
            if "-" in stripped and stripped.upper() == stripped.split()[0]:
                tickets.append(stripped.split()[0])
            elif "-" in stripped and not stripped.startswith("  "):
                tickets.append(stripped)
        if tickets:
            return sorted(set(tickets))
        simple = [line for line in lines if "-" in line and not line.strip().startswith("Add")]
        if simple:
            return sorted(set(simple))

    tickets = []
    for line in _parse_lines(data.get("change_rows", "")):
        if "|" not in line:
            continue
        parts = [part.strip() for part in line.split("|")]
        if parts and parts[0] and "-" in parts[0]:
            tickets.append(parts[0])
        elif len(parts) > 1 and parts[1]:
            tickets.append(parts[1])
    return sorted(set(tickets))


def _write_testing_evidence_section(doc, data: dict[str, Any]) -> None:
    """Testing Evidence: clients (deployment/config screenshots) then ticket list with placeholders."""
    from generators.release_notes_template_defaults import (
        TESTING_EVIDENCE_CLIENT_PLACEHOLDER,
        TESTING_EVIDENCE_CLIENTS_HEADING,
        TESTING_EVIDENCE_HEADING,
        TESTING_EVIDENCE_SUBHEADING,
        TESTING_EVIDENCE_TICKET_PLACEHOLDER,
        TESTING_EVIDENCE_TICKETS_LIST_HEADING,
        build_testing_evidence_client_names,
    )

    extracted = data.get("extracted") or []
    deployment_rows = data.get("deployment_rows", "")
    evidence_tickets = _testing_evidence_ticket_lines(data)
    client_names = build_testing_evidence_client_names(extracted, deployment_rows)

    h = doc.add_paragraph(style="Heading 2")
    h_run = h.add_run(f"{TESTING_EVIDENCE_HEADING}\u00a0")
    _style_run(h_run, bold=True)
    _format_section_paragraph(h, shaded=True, spacing="shaded_section")
    _set_paragraph_page_break_before(h)

    sub = doc.add_paragraph()
    sub_run = sub.add_run(TESTING_EVIDENCE_SUBHEADING)
    _style_run(sub_run, bold=True)
    _bind_block(h, sub)

    if client_names:
        clients_heading = doc.add_paragraph()
        clients_run = clients_heading.add_run(TESTING_EVIDENCE_CLIENTS_HEADING)
        _style_run(clients_run, bold=True)
        _format_section_paragraph(clients_heading, spacing="testing_subheading")
        for client in client_names:
            client_para = doc.add_paragraph()
            client_run = client_para.add_run(client)
            _style_run(client_run, bold=True)
            placeholder = doc.add_paragraph()
            ph_run = placeholder.add_run(TESTING_EVIDENCE_CLIENT_PLACEHOLDER)
            _style_run(ph_run, italic=True)
            _bind_block(client_para, placeholder)
            _format_section_paragraph(placeholder, spacing="testing_client_block")

    tickets_heading = doc.add_paragraph()
    tickets_run = tickets_heading.add_run(TESTING_EVIDENCE_TICKETS_LIST_HEADING)
    _style_run(tickets_run, bold=True)
    _format_section_paragraph(tickets_heading, spacing="testing_subheading")
    for idx, ticket in enumerate(evidence_tickets):
        ticket_para = doc.add_paragraph()
        ticket_run = ticket_para.add_run(ticket)
        _style_run(ticket_run)
        placeholder = doc.add_paragraph()
        ph_run = placeholder.add_run(TESTING_EVIDENCE_TICKET_PLACEHOLDER)
        _style_run(ph_run, italic=True)
        if idx == 0:
            _bind_block(tickets_heading, ticket_para)
        _bind_block(ticket_para, placeholder)
        _format_section_paragraph(placeholder, spacing="testing_ticket_block")


def _screenshot_ticket_lines(data: dict[str, Any]) -> list[str]:
    """Backward-compatible alias for :func:`_testing_evidence_ticket_lines`."""
    return _testing_evidence_ticket_lines(data)


def _parse_excluded_rows(raw: str) -> list[list[str]]:
    """Parse excluded ticket lines into 8-column rows (pipe-delimited or ticket-only)."""
    rows: list[list[str]] = []
    for line in _parse_lines(raw):
        if "|" in line:
            cols = [c.strip() for c in line.split("|", len(EXCLUDED_HEADERS) - 1)]
        else:
            cols = ["", line.strip(), "", "", "", "", "", ""]
        while len(cols) < len(EXCLUDED_HEADERS):
            cols.append("")
        rows.append(cols[: len(EXCLUDED_HEADERS)])
    return rows


def _write_excluded_header_cell(cell, header: str) -> None:
    _reset_cell_paragraph(cell)
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(header)
    _style_run(run, size=BODY_SIZE, color_hex=EXCLUDED_HEADER_TEXT)
    _set_cell_shading(cell, EXCLUDED_HEADER_FILL)
    _add_cell_borders(cell)
    _set_cell_no_wrap(cell)


def _write_excluded_data_cell(cell, value: str) -> None:
    text = (value or "").strip()
    _reset_cell_paragraph(cell)
    if text:
        run = cell.paragraphs[0].add_run(text)
        _style_run(run, size=BODY_SIZE)
    _set_cell_shading(cell, EXCLUDED_BODY_FILL)
    _add_cell_borders(cell)


def _write_excluded_table(table, excluded_rows: list[list[str]]) -> None:
    _apply_reference_table_style(table, EXCLUDED_TABLE_STYLE)
    _add_table_borders(table)

    header_row = table.rows[0]
    for col_idx, header in enumerate(EXCLUDED_HEADERS):
        _write_excluded_header_cell(header_row.cells[col_idx], header)

    body_rows = excluded_rows or [[]]
    for row_idx, row_data in enumerate(body_rows):
        row = table.rows[row_idx + 1]
        for col_idx in range(len(EXCLUDED_HEADERS)):
            value = row_data[col_idx] if col_idx < len(row_data) else ""
            _write_excluded_data_cell(row.cells[col_idx], value)

    _apply_table_column_widths(
        table,
        EXCLUDED_COLUMN_WIDTHS,
        total_width=EXCLUDED_TABLE_WIDTH,
        fixed_layout=True,
    )
    _prevent_table_row_splits(table)


def _change_table_has_analyze_rows(change_rows: list[list[str]], data: dict[str, Any]) -> bool:
    extracted = data.get("extracted") or []
    if any(entry.get("type") == "Analyze" for entry in extracted):
        return True
    for row in change_rows:
        if len(row) >= 12:
            if (row[5] or "").lower() == "analyze" or (row[6] or "").lower() == "standard":
                return True
    return False


def generate_release_notes_docx(data: dict[str, Any]) -> bytes:
    """Generate a release notes Word document matching the reference template."""
    doc = Document()
    _apply_reference_margins(doc)
    _configure_reference_document_theme(doc)

    style = doc.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = BODY_SIZE
    style.paragraph_format.space_after = Pt(0)

    version = data.get("version", "")
    if version and not str(version).lower().startswith("v"):
        version = f"v{version}"

    release_date = data.get("release_date", "")

    # ── Title & metadata ──────────────────────────────────────────
    title = doc.add_paragraph(style="Heading 1")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Release Notes ")
    _style_run(title_run, size=BODY_SIZE)

    if version:
        version_para = doc.add_paragraph()
        version_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for part in _split_version_run_parts(version):
            version_run = version_para.add_run(part)
            _style_run(version_run, bold=True, color_hex=VERSION_TEXT_COLOR)

    doc.add_paragraph("")

    if release_date:
        _add_label_value_line(doc, "Release Date", release_date)
    _add_label_value_line(doc, "Development Team", data.get("dev_team", ""))
    _add_label_value_line(doc, "Release Owner", data.get("release_owner", ""))
    _add_label_value_line(doc, "Senior Manager", data.get("senior_mgr", ""))
    _add_label_value_line(doc, "Director", data.get("director", ""))
    _add_label_value_line(doc, METADATA_GM_LABEL, data.get("svp", ""))
    _add_label_value_line(doc, "Environment", data.get("environment", "Production"))

    doc.add_paragraph("")

    # ── Release Summary ───────────────────────────────────────────
    summary_p1 = data.get("summary_p1", "")
    summary_para = doc.add_paragraph()
    summary_label = summary_para.add_run("Release Summary ")
    _style_run(summary_label)
    if summary_p1:
        summary_para.add_run("\n")
        summary_body = summary_para.add_run(summary_p1)
        _style_run(summary_body)
    _format_section_paragraph(summary_para, shaded=True, spacing="shaded_section")

    highlight_tables = resolve_highlight_summary_tables(data)
    if highlight_tables:
        for block in highlight_tables:
            _add_summary_table(doc, block.headers, block.rows)
            doc.add_paragraph("")

    # ── Key Highlights ────────────────────────────────────────────
    highlights = _parse_key_highlights(data.get("summary_p2", ""))
    if not highlights and data.get("extracted"):
        auto_lines = build_key_highlight_lines(data.get("extracted") or [])
        highlights = _parse_key_highlights("\n".join(auto_lines))
    if highlights:
        kh = doc.add_paragraph()
        kh_run = kh.add_run("Key Highlights")
        _style_run(kh_run, bold=True)
        _format_section_paragraph(kh, shaded=True, spacing="shaded_section")
        for prefix, rest in highlights:
            bullet = doc.add_paragraph(style=LIST_STYLE)
            if prefix:
                bold_run = bullet.add_run(f"{prefix}:")
                _style_run(bold_run, bold=True)
                if rest:
                    bullet.add_run(f" {rest}")
                    _style_run(bullet.runs[-1])
            elif rest:
                body = bullet.add_run(rest)
                _style_run(body)
            _format_section_paragraph(bullet, shaded=True, spacing="shaded_section")

    # ── Reports / Dashboards Delivered ────────────────────────────
    delivered_rows = _resolve_delivered_reports_rows(data)
    change_rows = _parse_change_rows(data.get("change_rows", ""))

    if delivered_rows:
        delivered_heading = doc.add_paragraph()
        delivered_run = delivered_heading.add_run("Reports / Dashboards Delivered")
        _style_run(delivered_run, bold=True)
        _format_section_paragraph(delivered_heading, shaded=True, spacing="shaded_section")
        _set_paragraph_keep_with_next(delivered_heading)
        _add_report_inventory_table(
            doc,
            delivered_rows,
            headers=DELIVERED_REPORTS_HEADERS,
            column_widths=DELIVERED_REPORTS_COLUMN_WIDTHS,
            indent_twips=DEPLOYMENT_TBL_IND_TWIPS,
        )
        doc.add_paragraph("")

    # ── Deployment tracking table ─────────────────────────────────
    if delivered_rows:
        _add_reference_gap_paragraphs(doc)
    deployment_rows = _parse_pipe_rows(data.get("deployment_rows", ""), min_cols=4)
    table = doc.add_table(rows=1 + max(len(deployment_rows), 0), cols=4)
    _write_deployment_table(table, deployment_rows)

    _add_post_deployment_table_gaps(doc)

    # ── Ticket Details ────────────────────────────────────────────
    heading = doc.add_paragraph()
    cd_run = heading.add_run("Ticket Details:")
    _style_run(cd_run, bold=True, color_hex=VERSION_TEXT_COLOR)
    _format_section_paragraph(heading, spacing="change_details_heading")
    _set_paragraph_keep_with_next(heading)

    _add_pre_change_table_gaps(doc)

    table = doc.add_table(rows=1 + len(change_rows), cols=len(CHANGE_HEADERS))
    _write_change_table(table, change_rows)

    _add_post_change_table_gaps(doc)

    # ── Testing Evidence ──────────────────────────────────────────
    _write_testing_evidence_section(doc, data)
    _add_doc_spacer(doc, "post_testing_gap")

    # ── Tickets excluded ──────────────────────────────────────────
    excluded_raw = data.get("excluded_tickets", "").strip()
    excluded_para = doc.add_paragraph()
    excluded_run = excluded_para.add_run("Tickets Excluded from Deployment")
    _style_run(excluded_run, bold=True)
    _format_section_paragraph(excluded_para, spacing="excluded_heading")
    _set_paragraph_keep_with_next(excluded_para)

    excluded_rows = _parse_excluded_rows(excluded_raw)
    table = doc.add_table(
        rows=1 + max(len(excluded_rows), 1),
        cols=len(EXCLUDED_HEADERS),
    )
    _write_excluded_table(table, excluded_rows)

    # ── Notes ─────────────────────────────────────────────────────
    notes = data.get("notes", "").strip()
    notes_heading = doc.add_paragraph()
    notes_run = notes_heading.add_run("Notes:")
    _style_run(notes_run, bold=True)
    if notes:
        for line in _parse_lines(notes):
            para = doc.add_paragraph()
            line_run = para.add_run(line)
            _style_run(line_run)

    buf = io.BytesIO()
    doc.core_properties.keywords = GENERATOR_BUILD_ID
    doc.core_properties.comments = f"jira-release-tool generator {GENERATOR_BUILD_ID}"
    doc.save(buf)
    return patch_docx_grid_table4_accent1_style(buf.getvalue())
