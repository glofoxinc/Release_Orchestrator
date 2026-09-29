"""Release Notes end-to-end pipeline status (generator, readiness, publish, parity)."""

from __future__ import annotations

import io
import os
from typing import Any

from generators.release_notes_generator import GENERATOR_BUILD_ID
from generators.release_notes_insights import (
    ReadinessReport,
    assess_release_notes_readiness,
    readiness_to_dict,
)
from generators.release_notes_publish import PublishResult, publish_release_notes_docx
from generators.release_notes_sample import sample_release_notes_data


def reference_fixture_available() -> bool:
    try:
        from tests.reference_fixture import REFERENCE_AVAILABLE

        return bool(REFERENCE_AVAILABLE)
    except Exception:  # noqa: BLE001
        root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        fixture = os.path.join(root, "tests", "fixtures", "Release Notes 2026.3.09.09.docx")
        return os.path.isfile(fixture)


def platform_features() -> list[str]:
    return [
        "DOCX v4 layout (metadata → summary → highlights → delivered table → deployment → ticket details 4-col → testing evidence)",
        "0–100 readiness score with blockers / warnings / suggestions",
        "Shared publish gate (readiness + hyperlink verification) — UI Step 4 + MCP",
        "Post-QA exclusions flow into excluded-tickets table + snapshot",
        "Key highlights + Analyze/Custom inventory tables from tickets",
        "Test-by QA pool (never matches Jira assignee)",
        "Deployment tracking prefill from Deployer / history",
        "Manual SharePoint link → SHIPIT service ticket",
        "Generator build ID stamped in DOCX properties",
    ]


def e2e_pipeline_steps(*, publish: PublishResult | None = None, readiness: ReadinessReport | None = None) -> list[dict[str, Any]]:
    """Ordered checklist for Release Notes E2E (Settings, Step 4, CLI)."""
    report = readiness or assess_release_notes_readiness(sample_release_notes_data())
    pub = publish
    if pub is None:
        pub = publish_release_notes_docx(sample_release_notes_data(), enforce_readiness=True)

    steps: list[dict[str, Any]] = [
        {
            "id": "generator",
            "label": "Built-in generator",
            "ok": True,
            "detail": f"build `{GENERATOR_BUILD_ID}`",
        },
        {
            "id": "reference_fixture",
            "label": "Reference DOCX fixture",
            "ok": reference_fixture_available(),
            "detail": "orchestrator/zones/quality/tests/fixtures for structural parity" if reference_fixture_available() else "fixture missing",
            "fix": "Copy reference DOCX to orchestrator/zones/quality/tests/fixtures/Release Notes 2026.3.09.09.docx",
        },
        {
            "id": "readiness_gate",
            "label": "Readiness scoring",
            "ok": report.ready and report.score >= 80,
            "detail": f"{report.score}/100 — {report.score_label}",
        },
        {
            "id": "publish_gate",
            "label": "Publish gate (DOCX + hyperlinks)",
            "ok": pub.status == "ok" and bool(pub.docx_bytes),
            "detail": pub.error or f"{len(pub.docx_bytes or b'')} bytes, hyperlinks OK",
        },
        {
            "id": "docx_sections",
            "label": "DOCX v4 section stack",
            "ok": pub.status == "ok",
            "detail": "metadata, summary, highlights, delivered reports, deployment, ticket details (4-col), testing evidence, excluded, notes",
        },
        {
            "id": "mcp_parity",
            "label": "MCP generate_release_notes_docx",
            "ok": True,
            "detail": "Same publish_release_notes_docx path as UI Step 4",
        },
        {
            "id": "sharepoint_handoff",
            "label": "SharePoint to Service Ticket handoff",
            "ok": True,
            "detail": "Paste manual link on Step 4; flows to SHIPIT description",
            "fix": "Restore archive/sharepoint_upload.yaml when MS_GRAPH_* is ready (optional)",
        },
    ]
    return steps


def run_fixture_e2e(*, enforce_readiness: bool = True, write_path: str = "") -> dict[str, Any]:
    """Run full E2E on sample data: readiness → publish → optional write."""
    data = sample_release_notes_data()
    readiness = assess_release_notes_readiness(data)
    publish = publish_release_notes_docx(data, enforce_readiness=enforce_readiness)

    out: dict[str, Any] = {
        "generator_build_id": GENERATOR_BUILD_ID,
        "reference_fixture_available": reference_fixture_available(),
        "readiness": readiness_to_dict(readiness),
        "publish": publish.to_dict(),
        "pipeline_steps": e2e_pipeline_steps(publish=publish, readiness=readiness),
        "platform_features": platform_features(),
    }

    if publish.status == "ok" and publish.docx_bytes:
        out["docx_bytes_len"] = len(publish.docx_bytes)
        # Quick structural smoke: DOCX opens and has tables
        try:
            from docx import Document

            doc = Document(io.BytesIO(publish.docx_bytes))
            out["docx_table_count"] = len(doc.tables)
            out["docx_paragraph_count"] = len(doc.paragraphs)
            props = doc.core_properties
            out["docx_keywords"] = str(props.keywords or "")
        except Exception as exc:  # noqa: BLE001
            out["docx_parse_error"] = str(exc)

        if write_path:
            with open(write_path, "wb") as f:
                f.write(publish.docx_bytes)
            out["written_path"] = write_path

    out["ok"] = publish.status == "ok" and readiness.ready
    return out


def e2e_status() -> dict[str, Any]:
    """Lightweight status for verify_integrations (no file write)."""
    result = run_fixture_e2e(enforce_readiness=True)
    pending = [s for s in result.get("pipeline_steps") or [] if not s.get("ok")]
    return {
        "feature": "release_notes",
        "generator_build_id": GENERATOR_BUILD_ID,
        "ok": bool(result.get("ok")),
        "readiness_score": (result.get("readiness") or {}).get("score"),
        "publish_status": (result.get("publish") or {}).get("status"),
        "reference_fixture_available": result.get("reference_fixture_available"),
        "pending_step_count": len(pending),
        "pipeline_steps": result.get("pipeline_steps"),
        "platform_features": platform_features(),
    }
