"""Render PNG page previews from the built-in Release Notes generator (v4 format).

Usage:
    python scripts/render_release_notes_samples.py
    python scripts/render_release_notes_samples.py --out docs/samples/release-notes
"""

from __future__ import annotations

import argparse
import io
import os
import sys
import textwrap
from typing import Any

_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)

from docx import Document
from PIL import Image, ImageDraw, ImageFont

from generators.release_notes_generator import GENERATOR_BUILD_ID, generate_release_notes_docx
from generators.release_notes_sample import sample_release_notes_data
from generators.release_notes_report_tables import build_delivered_reports_rows
from generators.release_notes_template_defaults import (
    DELIVERED_REPORTS_HEADERS,
    TESTING_EVIDENCE_HEADING,
    TESTING_EVIDENCE_SUBHEADING,
    TITLE_COLOR,
)

# US Letter @ ~96 DPI for readable previews
PAGE_W, PAGE_H = 816, 1056
MARGIN = 72
CONTENT_W = PAGE_W - 2 * MARGIN

TITLE_RGB = tuple(int(TITLE_COLOR[i : i + 2], 16) for i in (0, 2, 4))
HEADER_FILL = (68, 114, 196)  # Grid Accent 1 approx
SECTION_FILL = (247, 247, 247)
BORDER = (180, 180, 180)
LINK_BLUE = (5, 99, 193)
TEXT = (30, 30, 30)
MUTED = (100, 100, 100)


def _font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = (
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/segoeui.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    )
    for path in candidates:
        if os.path.isfile(path):
            try:
                return ImageFont.truetype(path, size)
            except OSError:
                continue
    return ImageFont.load_default()


def _wrap(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont, width: int) -> list[str]:
    if not text.strip():
        return []
    avg = max(draw.textlength("abcdefghijklmnopqrstuvwxyz", font=font) / 26, 6)
    chars = max(20, int(width / avg))
    return textwrap.wrap(text, width=chars)


def _new_page() -> tuple[Image.Image, ImageDraw.ImageDraw, int]:
    img = Image.new("RGB", (PAGE_W, PAGE_H), "white")
    draw = ImageDraw.Draw(img)
    return img, draw, MARGIN


def _ensure_space(img, draw, y, needed, pages, page_num):
    if y + needed <= PAGE_H - MARGIN:
        return img, draw, y, page_num
    pages.append(img)
    img, draw, y = _new_page()
    return img, draw, y, page_num + 1


def _draw_centered(draw, y, text, font, fill=TEXT):
    w = draw.textlength(text, font=font)
    draw.text(((PAGE_W - w) / 2, y), text, fill=fill, font=font)
    return y + font.size + 8


def _draw_section_bar(draw, y, label, font_bold):
    draw.rectangle([MARGIN, y, PAGE_W - MARGIN, y + 22], fill=SECTION_FILL, outline=BORDER)
    draw.text((MARGIN + 8, y + 3), label, fill=TEXT, font=font_bold)
    return y + 30


def _draw_meta_lines(draw, y, data: dict[str, Any], font, font_bold):
    rows = [
        ("Release Date", data.get("release_date", "")),
        ("Development Team", data.get("dev_team", "")),
        ("Release Owner", data.get("release_owner", "")),
        ("Senior Manager", data.get("senior_mgr", "")),
        ("Director", data.get("director", "")),
        ("GM", data.get("svp", "")),
        ("Environment", data.get("environment", "")),
    ]
    for label, value in rows:
        draw.text((MARGIN, y), f"{label}:", fill=TEXT, font=font_bold)
        draw.text((MARGIN + 200, y), str(value), fill=TEXT, font=font)
        y += 18
    return y + 8


def _draw_table(
    draw,
    y,
    headers: list[str],
    rows: list[list[str]],
    col_widths: list[int],
    font,
    font_bold,
    *,
    link_cols: set[int] | None = None,
) -> int:
    link_cols = link_cols or set()
    x0 = MARGIN
    row_h = 22
    header_h = 24

    def cell_x(col: int) -> int:
        return x0 + sum(col_widths[:col])

    # Header
    draw.rectangle([x0, y, x0 + sum(col_widths), y + header_h], fill=HEADER_FILL)
    for i, h in enumerate(headers):
        draw.text((cell_x(i) + 4, y + 4), h[:18], fill="white", font=font_bold)
    y += header_h

    for row in rows:
        draw.rectangle([x0, y, x0 + sum(col_widths), y + row_h], outline=BORDER)
        for i, cell in enumerate(row[: len(headers)]):
            color = LINK_BLUE if i in link_cols and str(cell).startswith("http") else TEXT
            label = str(cell)
            if len(label) > 22:
                label = label[:19] + "..."
            draw.text((cell_x(i) + 4, y + 4), label, fill=color, font=font)
        y += row_h
    return y + 12


def render_preview_pages(data: dict[str, Any]) -> list[Image.Image]:
    font = _font(11)
    font_bold = _font(11, bold=True)
    font_title = _font(20, bold=True)
    font_version = _font(13, bold=True)

    pages: list[Image.Image] = []
    img, draw, y = _new_page()
    page_num = 1

    version = data.get("version", "")
    if version and not str(version).lower().startswith("v"):
        version = f"v{version}"

    y = _draw_centered(draw, y, "Release Notes", font_title, fill=TITLE_RGB)
    y = _draw_centered(draw, y, f"Version: {version}", font_version, fill=TEXT)
    y += 10
    y = _draw_meta_lines(draw, y, data, font, font_bold)

    y = _draw_section_bar(draw, y, "Release Summary", font_bold)
    for line in _wrap(draw, data.get("summary_p1", ""), font, CONTENT_W - 16):
        draw.text((MARGIN + 8, y), line, fill=TEXT, font=font)
        y += 16
    y += 6

    y = _draw_section_bar(draw, y, "Key Highlights", font_bold)
    for line in (data.get("summary_p2") or "").splitlines():
        if line.strip():
            draw.text((MARGIN + 8, y), f"• {line.strip()}", fill=TEXT, font=font)
            y += 16
    y += 8

    y = _draw_section_bar(draw, y, "Reports / Dashboards Delivered", font_bold)
    delivered = build_delivered_reports_rows(data.get("extracted") or [])
    if not delivered and data.get("delivered_report_rows"):
        delivered = [
            [part.strip() for part in line.split("|", 2)]
            for line in str(data.get("delivered_report_rows")).splitlines()
            if line.strip()
        ]
    dr_headers = list(DELIVERED_REPORTS_HEADERS)
    dr_widths = [120, 220, CONTENT_W - 340]
    y = _draw_table(draw, y, dr_headers, delivered[:6], dr_widths, font, font_bold)

    img, draw, y, page_num = _ensure_space(img, draw, y, 200, pages, page_num)
    y = _draw_section_bar(draw, y, "Deployment Tracking", font_bold)
    dep_headers = ["Client", "Deployment Owner", "Status", "Deployment ID & Time"]
    dep_rows = [r.split("|") for r in (data.get("deployment_rows") or "").splitlines() if r.strip()]
    dep_rows = [[c.strip() for c in row] + [""] * (4 - len(row)) for row in dep_rows]
    dw = [120, 140, 100, CONTENT_W - 360]
    y = _draw_table(draw, y, dep_headers, dep_rows[:3], dw, font, font_bold)

    y += 4
    draw.text((MARGIN, y), "Ticket Details:", fill=TEXT, font=font_bold)
    y += 22
    ch_headers = ["Ticket", "Ticket Link", "Test", "Status"]
    ch_rows = []
    for line in (data.get("change_rows") or "").splitlines():
        if not line.strip():
            continue
        parts = [part.strip() for part in line.split("|")]
        if len(parts) >= 4:
            ch_rows.append(parts[:4])
    cw = [80, 260, 80, CONTENT_W - 420]
    y = _draw_table(draw, y, ch_headers, ch_rows, cw, font, font_bold, link_cols={1})

    img, draw, y, page_num = _ensure_space(img, draw, y, 160, pages, page_num)
    pages.append(img)
    img, draw, y = _new_page()
    page_num += 1
    y = _draw_section_bar(draw, y, TESTING_EVIDENCE_HEADING.rstrip(":"), font_bold)
    draw.text((MARGIN + 8, y), TESTING_EVIDENCE_SUBHEADING, fill=MUTED, font=font_bold)
    y += 20
    evidence = data.get("testing_evidence_tickets") or data.get("screenshot_names") or ""
    for ticket in str(evidence).splitlines():
        if ticket.strip():
            draw.text((MARGIN + 8, y), ticket.strip(), fill=TEXT, font=font)
            y += 16

    draw.text((MARGIN, PAGE_H - 40), f"Generator: {GENERATOR_BUILD_ID}", fill=MUTED, font=font)
    draw.text((MARGIN, PAGE_H - 24), "Preview — Insights Customize BI Team · Production release sample", fill=MUTED, font=font)
    pages.append(img)
    return pages


def main() -> int:
    parser = argparse.ArgumentParser(description="Render Release Notes sample PNG previews")
    parser.add_argument(
        "--out",
        default=os.path.join(_ROOT, "docs", "samples", "release-notes"),
        help="Output directory for PNG + DOCX",
    )
    args = parser.parse_args()
    os.makedirs(args.out, exist_ok=True)

    data = sample_release_notes_data()
    data["dev_team"] = "Insights Customize BI Team"
    data["version"] = "v2026.3.08.12"
    data["release_date"] = "12-August-2026"

    docx_bytes = generate_release_notes_docx(data)
    docx_path = os.path.join(args.out, "Release Notes 2026.3.08.12 SAMPLE.docx")
    with open(docx_path, "wb") as f:
        f.write(docx_bytes)

    # Structural smoke from real DOCX
    doc = Document(io.BytesIO(docx_bytes))
    meta_path = os.path.join(args.out, "sample-meta.txt")
    with open(meta_path, "w", encoding="utf-8") as f:
        f.write(f"generator_build={GENERATOR_BUILD_ID}\n")
        f.write(f"paragraphs={len(doc.paragraphs)}\n")
        f.write(f"tables={len(doc.tables)}\n")
        f.write(f"docx_bytes={len(docx_bytes)}\n")

    pages = render_preview_pages(data)
    png_paths: list[str] = []
    for i, page in enumerate(pages, start=1):
        path = os.path.join(args.out, f"release-notes-sample-page-{i}.png")
        page.save(path, "PNG", optimize=True)
        png_paths.append(path)

    print(f"Wrote {docx_path}")
    for p in png_paths:
        print(f"Wrote {p}")
    print(f"Pages: {len(png_paths)} | build {GENERATOR_BUILD_ID}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
