"""Validate generated DOCX structure against the reference template."""

from __future__ import annotations

import io
import os
import sys

import pytest
from docx import Document

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from generators.release_notes_generator import generate_release_notes_docx
from generators.release_notes_template_defaults import CHANGE_HEADERS
from tests.reference_fixture import REFERENCE_AVAILABLE, REFERENCE_PATH
from docx.oxml.ns import qn


def _reference_like_data() -> dict:
    from generators.release_notes_sample import sample_release_notes_data

    return sample_release_notes_data()


@pytest.mark.skipif(not REFERENCE_AVAILABLE, reason="Reference DOCX not available")
class TestReferenceDocPresent:
    def test_reference_has_expected_sections(self):
        from docx import Document

        doc = Document(REFERENCE_PATH)
        text = "\n".join(p.text for p in doc.paragraphs)
        for phrase in (
            "Release Notes",
            "Key Highlights",
            "Reports / Dashboards Delivered",
            "Ticket Details:",
            "Testing Evidence",
            "Tickets Excluded from Deployment",
            "Notes:",
        ):
            assert phrase in text


class TestGeneratedMatchesReferenceStructure:
    def test_section_order_and_tables(self):
        from docx import Document

        result = generate_release_notes_docx(_reference_like_data())
        doc = Document(io.BytesIO(result))
        text = "\n".join(p.text for p in doc.paragraphs)

        assert "Release Notes" in text
        assert "Version: v2026.3.09.09" in text
        assert "Release Date" in text and "15-July-2026" in text
        assert "Key Highlights" in text
        assert "Reports / Dashboards Delivered" in text
        assert "Ticket Details:" in text
        assert "Sanity Checklist After Production Deployment:" not in text
        assert "Testing Evidence:" in text
        assert "Tickets Excluded from Deployment" in text
        assert "Notes:" in text

        assert len(doc.tables) >= 4
        change_table = next(
            t for t in doc.tables if len(t.columns) == 5 and any(
                c.text.strip() == "Ticket" for c in t.rows[0].cells
            )
        )
        headers = [c.text.strip() for c in change_table.rows[0].cells]
        assert headers == CHANGE_HEADERS

        deploy_table = next(
            t for t in doc.tables if any("Deployment Owner" in c.text for c in t.rows[0].cells)
        )
        assert "Client" in deploy_table.rows[0].cells[0].text

    def test_delivered_and_deployment_tables_keep_alignment_and_row_integrity(self):
        result = generate_release_notes_docx(_reference_like_data())
        doc = Document(io.BytesIO(result))
        delivered = next(
            t for t in doc.tables if any("Dashboard / Report" in c.text for c in t.rows[0].cells)
        )
        deployment = next(
            t for t in doc.tables if any("Deployment Owner" in c.text for c in t.rows[0].cells)
        )

        delivered_ind = delivered._tbl.tblPr.find(qn("w:tblInd"))
        deployment_ind = deployment._tbl.tblPr.find(qn("w:tblInd"))
        assert delivered_ind is not None
        assert deployment_ind is not None
        assert delivered_ind.get(qn("w:w")) == deployment_ind.get(qn("w:w"))

        delivered_header_pr = delivered.rows[0]._tr.get_or_add_trPr()
        deployment_header_pr = deployment.rows[0]._tr.get_or_add_trPr()
        assert delivered_header_pr.find(qn("w:tblHeader")) is not None
        assert deployment_header_pr.find(qn("w:tblHeader")) is not None

        for row in delivered.rows:
            tr_pr = row._tr.get_or_add_trPr()
            assert tr_pr.find(qn("w:cantSplit")) is not None
        for row in deployment.rows:
            tr_pr = row._tr.get_or_add_trPr()
            assert tr_pr.find(qn("w:cantSplit")) is not None
