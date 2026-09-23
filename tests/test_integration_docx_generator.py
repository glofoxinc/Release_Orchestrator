"""Integration tests for the in-repo DOCX generator (generators/release_notes_generator.py)."""

from __future__ import annotations

import io
import os
import sys

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from generators.release_notes_generator import generate_release_notes_docx


def _minimal_data() -> dict:
    return {"version": "1.0.0"}


def _full_data() -> dict:
    change_lines = "\n".join(
        f"{i}|PROJ-{100+i}|PROJ-{100+i}|PARENT-{i}|https://jira.example.com/PROJ-{100+i}"
        f"|Bug|ClientA|Dashboard|Fix for item {i}|QA Tester||Done"
        for i in range(1, 7)
    )
    return {
        "version": "v2.5.0",
        "release_date": "15-July-2026",
        "dev_team": "Platform Engineering",
        "release_owner": "Jane Doe",
        "senior_mgr": "John Smith",
        "director": "Alice B.",
        "svp": "Bob C.",
        "environment": "Production",
        "summary_p1": "This release includes critical bug fixes and new dashboard features.",
        "summary_p2": "Standard (Analyze): KPI tile updates\nLift: New paginated report",
        "analyze_clients": "Analyze clients",
        "custom_clients": "Acme Corp, Globex",
        "std_summary": "Standard client deployments remain unchanged.",
        "dashboard_items": "Analyze- Revenue Dashboard\nLift - Ops Monitor",
        "change_rows": change_lines,
        "deployment_rows": "Analyze|Navya||\nLift|Chanti||",
        "sanity_checklist": "Report loads successfully – Yes",
        "screenshot_names": "PIC-5730\nAN-5328",
    }


class TestGenerateMinimalData:
    """Verify the generator handles the bare-minimum input."""

    def test_returns_bytes(self):
        result = generate_release_notes_docx(_minimal_data())
        assert isinstance(result, bytes)

    def test_non_empty(self):
        result = generate_release_notes_docx(_minimal_data())
        assert len(result) > 0

    def test_zip_magic_bytes(self):
        result = generate_release_notes_docx(_minimal_data())
        assert result[:2] == b"PK", "DOCX must start with PK (ZIP magic bytes)"


class TestGenerateFullData:
    """Verify a richly-populated data dict produces a valid DOCX."""

    def test_returns_non_empty_bytes(self):
        result = generate_release_notes_docx(_full_data())
        assert isinstance(result, bytes) and len(result) > 0

    def test_zip_magic_bytes(self):
        result = generate_release_notes_docx(_full_data())
        assert result[:2] == b"PK"

    def test_loadable_by_python_docx(self):
        from docx import Document

        result = generate_release_notes_docx(_full_data())
        doc = Document(io.BytesIO(result))
        assert doc is not None

    def test_contains_title_text(self):
        from docx import Document

        result = generate_release_notes_docx(_full_data())
        doc = Document(io.BytesIO(result))
        all_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Release Notes" in all_text

    def test_contains_version(self):
        from docx import Document

        data = _full_data()
        result = generate_release_notes_docx(data)
        doc = Document(io.BytesIO(result))
        all_text = "\n".join(p.text for p in doc.paragraphs)
        assert data["version"] in all_text

    def test_contains_summary_section(self):
        from docx import Document

        result = generate_release_notes_docx(_full_data())
        doc = Document(io.BytesIO(result))
        all_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Release Summary" in all_text

    def test_contains_change_details_table(self):
        from docx import Document

        result = generate_release_notes_docx(_full_data())
        doc = Document(io.BytesIO(result))
        assert len(doc.tables) >= 2, "Expected at least the metadata table and the change details table"

    def test_change_table_has_header_row(self):
        from docx import Document

        from generators.release_notes_template_defaults import CHANGE_HEADERS

        result = generate_release_notes_docx(_full_data())
        doc = Document(io.BytesIO(result))
        change_table = next(
            t for t in doc.tables
            if len(t.columns) == 5 and any(c.text.strip() == "Ticket" for c in t.rows[0].cells)
        )
        header_cells = [c.text.strip() for c in change_table.rows[0].cells]
        assert header_cells == CHANGE_HEADERS

    def test_contains_key_sections(self):
        from docx import Document

        result = generate_release_notes_docx(_full_data())
        doc = Document(io.BytesIO(result))
        all_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Key Highlights" in all_text
        assert "Sanity Checklist After Production Deployment" not in all_text
        assert "Testing Evidence:" in all_text

    def test_sanity_checklist_not_generated_by_default(self):
        from docx import Document

        result = generate_release_notes_docx(_minimal_data())
        doc = Document(io.BytesIO(result))
        all_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Sanity Checklist After Production Deployment" not in all_text

    def test_testing_evidence_section_lists_tickets_with_deployment_heading(self):
        from docx import Document

        from generators.release_notes_template_defaults import (
            TESTING_EVIDENCE_SUBHEADING,
            TESTING_EVIDENCE_TICKETS_LIST_HEADING,
            TESTING_EVIDENCE_TICKET_PLACEHOLDER,
        )

        result = generate_release_notes_docx(_full_data())
        doc = Document(io.BytesIO(result))
        all_text = "\n".join(p.text for p in doc.paragraphs)
        assert TESTING_EVIDENCE_SUBHEADING in all_text
        assert TESTING_EVIDENCE_TICKETS_LIST_HEADING in all_text
        assert TESTING_EVIDENCE_TICKET_PLACEHOLDER in all_text
        assert "PIC-5730" in all_text
        assert "AN-5328" in all_text
        assert "Analyze" in all_text or "Lift" in all_text

    def test_excluded_table_uses_reason_column_header(self):
        from docx import Document

        from generators.release_notes_template_defaults import EXCLUDED_HEADERS

        result = generate_release_notes_docx(_full_data())
        doc = Document(io.BytesIO(result))
        excluded = next(
            t for t in doc.tables
            if any("Ticket#" in c.text for c in t.rows[0].cells)
            and len(t.columns) == 8
            and t.rows[0].cells[0].text.strip() == "S.No"
        )
        headers = [c.text.strip() for c in excluded.rows[0].cells]
        assert headers == EXCLUDED_HEADERS
        assert "Reason" in headers
        assert "Description" not in headers


class TestSpecialCharacters:
    """Ensure special / unicode characters don't crash the generator."""

    @pytest.mark.parametrize(
        "summary_fragment",
        [
            "Fix <script>alert('xss')</script> in title",
            'Quotes "double" and \'single\'',
            "Accented: café, naïve, résumé",
            "Emoji: \U0001f680 rocket launch",
            "Ampersand & less-than < greater-than >",
            "Newline embedded\nin summary text",
        ],
    )
    def test_special_chars_in_summary(self, summary_fragment: str):
        data = _minimal_data()
        data["summary_p1"] = summary_fragment
        result = generate_release_notes_docx(data)
        assert result[:2] == b"PK"

    def test_special_chars_in_change_rows(self):
        data = _minimal_data()
        data["change_rows"] = (
            "1|PROJ-1|PROJ-1|P-1|http://x|Bug|C&C|R<1>|\"Fix\"|QA||Done"
        )
        result = generate_release_notes_docx(data)
        assert isinstance(result, bytes) and len(result) > 0

    def test_empty_strings_everywhere(self):
        data = {k: "" for k in _full_data()}
        data["version"] = ""
        result = generate_release_notes_docx(data)
        assert result[:2] == b"PK"
